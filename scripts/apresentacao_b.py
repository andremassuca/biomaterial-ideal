"""Apresentação no estilo B (15 out 2026): 10 minutos, 15 diapositivos principais e 11 de reserva (com o separador).

Uso:
    python scripts/apresentacao_b.py [--exportar] [--renders PASTA] [--sem-demo]

Saída: apresentacao/apresentacao_b.pptx; com --exportar (PowerPoint instalado), também apresentacao_b.pdf,
apresentacao_plano_B.pdf (o mesmo conteúdo, para a pen USB, sem depender dos vídeos) e
apresentacao/imagens_b/diapositivo_NN.png. O pptx exportado leva as letras embutidas.

Estilo B: capa e conclusão em azul-noite, conteúdo em branco, letra Inter, títulos-afirmação (o título diz a
conclusão do diapositivo), cartões claros e um só destaque quente. Imagens: renders do Blender
(scripts/renders_blender.py, variante transparente), gráficos de scripts/graficos.py e vídeos de
scripts/animacoes.py. Todos os números vêm de results/ e de data/.
"""
from __future__ import annotations

import argparse
import importlib.util
import sys
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "scripts"))
sys.path.insert(0, str(RAIZ / "src"))
import apresentacao as ap  # noqa: E402
import nomes_simples as ns  # noqa: E402
from figuras import curto  # noqa: E402

APRES = RAIZ / "apresentacao"
CACHE = APRES / "imagens_b" / "recursos"
DESTINO = APRES / "apresentacao_b.pptx"
RENDERS = [Path(r"partilha\renders\final"), Path(r"partilha\renders\rapida_8out")]

AUTOR = "André Oliveira Massuça"
ORIENTADOR = "Prof. Pedro Sampaio"
CABECALHO = "Próteses e Órgãos Artificiais · Licenciatura em Engenharia Biomédica · Universidade Lusófona"
AVISO = "Não é ferramenta clínica nem recomendação terapêutica."
N_PRINCIPAIS = 15  # diapositivos que contam para os 10 minutos (os seguintes são de reserva)
PALAVRAS_POR_MINUTO = 130  # ritmo de fala usado no projeto (R84)

ABERTURA = "abertura_anca_15s"  # vídeo de abertura (apresentacao/media/; original em partilha\renders)
URL_WEB = "andremassuca.github.io/biomaterial-ideal"  # versão web (docs/PUBLICACAO.md, ramo extensoes)
QR_WEB = APRES / "media" / "qr_versao_web.png"  # o mesmo QR do poster; aponta para https://<URL_WEB>/
# logótipo do projeto (versão para fundo escuro: a capa e a conclusão são escuras): o primeiro PNG que existir
# (o SVG não entra no PowerPoint pelo python-pptx); a cópia em apresentacao/media/ vem de partilha\logo\final
LOGO_PROJETO = [APRES / "media" / "logo_icone_fundo_escuro.png",
                Path(r"partilha\logo\final\logo_icone_fundo_escuro_1024.png")]
# vídeo da demonstração da app (reserva): o primeiro que existir
DEMO_APP = [RAIZ / "app-nativa" / "demo" / "demo_app_60s.mp4", Path(r"partilha\demo\demo_app_60s.mp4"),
            APRES / "media" / "demo_app.mp4"]

# tokens de cor do estilo B (os mesmos de scripts/design_direcoes.py)
NOITE = RGBColor(0x0B, 0x1B, 0x2E)
AZUL = RGBColor(0x05, 0x95, 0xD1)
AZUL_TEXTO = RGBColor(0x04, 0x77, 0xA8)
TEXTO = RGBColor(0x0F, 0x17, 0x20)
CINZA = RGBColor(0x55, 0x60, 0x6B)
CARTAO = RGBColor(0xF2, 0xF5, 0xF8)
LINHA = RGBColor(0xD5, 0xDD, 0xE5)
SUB_NOITE = RGBColor(0x9F, 0xB6, 0xC9)
QUENTE = RGBColor(0xB8, 0x5A, 0x12)  # laranja escuro: cerca de 4,6:1 com branco (antes #E8833A, 2,7:1)
QUENTE_FUNDO = RGBColor(0xFD, 0xF0, 0xE6)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
LETRA = "Inter"

DISPOSITIVOS = [  # (render, nome, o que tem de aguentar: o problema crítico indicado pelo docente)
    ("anca", "Haste femoral", "stress shielding e fadiga (problema acrescentado)"),
    ("articulacao", "Par articular", "desgaste e partículas"),
    ("stent_arteria", "Stent coronário", "resistência mecânica e biocompatibilidade"),
    ("scaffold", "Scaffold ósseo", "suporte celular e degradação controlada"),
    ("dentario", "Implante dentário", "integração com o osso e corrosão"),
]


# ---------------------------------------------------------------- dados (funções puras, testadas)

def pasta_renders(candidatas: list[Path] = RENDERS) -> Path:
    """Primeira pasta com os renders transparentes das cinco cenas (o 4K final, se já existir)."""
    for p in candidatas:
        if all((p / f"{c}_transparente.png").exists() for c, _, _ in DISPOSITIVOS):
            return p
    raise FileNotFoundError("sem renders transparentes das cinco cenas em " + ", ".join(map(str, candidatas)))


def confianca(numeros: ap.Numeros, leitura_muda_scaffold: bool,
              simples: bool = False) -> list[tuple[str, str, str, str]]:
    """Espectro de confiança: (dispositivo, material, nível, frase simples) para os cinco dispositivos.

    Em linguagem simples: cada escolha recebe um nível pelas regras do trabalho. Sólida: o vencedor ganha em
    pelo menos 99,9 % das repetições do Monte Carlo com os três métodos. Muito próximos ("Empate técnico" na versão 1,
    usada pelo poster): a margem ΔC fica abaixo de 0,01 (regra de empate da secção 2.8) em algum cenário. Depende da
    carga ("Depende da leitura" na versão 1): o vencedor do scaffold muda com a carga do defeito e com a leitura do
    problema crítico.
    Indicação (ou "Indicativo", com simples=True): casos semiquantitativos (notas de 1 a 5), em que o ranking só se
    lê como indicação. Com simples=True, os materiais levam os nomes simples (decisão de 8 out 2026); sem ele, ficam
    os nomes da versão 1 (usados pelo poster, que deixou de ser desenvolvido neste projeto)."""
    hc = numeros.mc_haste
    solida = min(hc.values()) >= 0.999 * numeros.n_iter
    stent_empate = min(numeros.dc_stent.values()) < 0.01
    if simples:
        mats = [ns.com_codigo("Ti-6Al-4V ELI"), "Platina-crómio (Pt-Cr) e cobalto-crómio (L605)",
                "Compósito PCL/β-TCP ou β-TCP poroso", ns.com_codigo(numeros.vencedor_par),
                ns.com_codigo(numeros.vencedor_dentario)]
        semi, f_par, f_dent = ("Indicativo", "a melhor nota no desgaste (notas de 1 a 5)",
                               "em 1.º em todas as versões, pela osteointegração (notas de 1 a 5)")
        f_stent, f_scaf = sem_prefixo(FRASE_STENT_CURTA), sem_prefixo(FRASE_SCAFFOLD_CURTA)
    else:
        mats = ["Ti-6Al-4V ELI", "Pt-Cr ou L605", "PCL/β-TCP", curto(numeros.vencedor_par),
                curto(numeros.vencedor_dentario)]
        semi, f_par, f_dent = ("Indicação", "a melhor nota no desgaste, numa escala de 1 a 5",
                               "em 1.º em todas as versões, pela osteointegração")
        f_stent = "praticamente equivalentes; o VIKOR não os separa"
        f_scaf = "muda para β-TCP se só contarem as células e a degradação"
    return [
        ("Haste femoral", mats[0], "Sólida" if solida else "Provável",
         "ganha com os três métodos em quase todas as repetições"),
        ("Stent coronário", mats[1], ("Muito próximos" if simples else "Empate técnico") if stent_empate else "Provável",
         f_stent),
        ("Scaffold ósseo", mats[2], ("Depende da carga" if simples else "Depende da leitura")
         if leitura_muda_scaffold else "Provável", f_scaf),
        ("Par articular", mats[3], semi, f_par),
        ("Implante dentário", mats[4], semi, f_dent),
    ]


def funil(results: Path = ap.RESULTS) -> dict[str, int]:
    """Números do funil de triagem (results/funil.csv, gerado por scripts/funil.py)."""
    t = pd.read_csv(results / "funil.csv")
    n = t.set_index("etapa")["n"]
    return {"mp": int(n["Materials Project (todos os compostos)"]),
            "nivel2": int(n["Módulo de Young dentro do intervalo"]),
            "elementos": int(n["Elementos avaliados"]), "normas": int(n["Normas ASTM de materiais para implantes"]),
            "base": int(n["Base verificada (materiais por componente)"]),
            "final": int(n["Recomendação (um material por dispositivo)"])}


def perfis_exemplo(results: Path = ap.RESULTS) -> list[tuple[str, str, str, str]]:
    """Perfis de referência em que o 1.º lugar (Borda) muda: (perfil, caso, antes, depois)."""
    p = pd.read_csv(results / "perfis_referencia.csv")
    saida = []
    for (_pid, desc, caso), g in p.groupby(["perfil", "descricao", "caso"], sort=False):
        antes = g.sort_values("pos_ref_Borda")["material"].iloc[0]
        depois = g.sort_values("pos_Borda")["material"].iloc[0]
        excl = str(g["excluidos"].iloc[0]) if pd.notna(g["excluidos"].iloc[0]) else ""
        if antes != depois or (excl and antes in excl):
            base = p[(p["caso"] == caso)].drop_duplicates("material")
            ref = base.sort_values("pos_ref_Borda")["material"].iloc[0]
            saida.append((desc, caso, ref, depois))
    return saida


def n_verificados(data: Path = ap.DATA) -> tuple[int, int]:
    """(linhas verificadas, linhas da base) em data/materiais.csv.

    Em linguagem simples: conta quantos valores da base têm um estado que começa por "Verificado" (fonte primária,
    fornecedor, derivado ou composição) e quantos valores há ao todo; os outros estão "A verificar"."""
    m = pd.read_csv(data / "materiais.csv")
    return int(m["estado"].astype(str).str.startswith("Verificado").sum()), len(m)


def mc_primeiro(caso: str, results: Path = ap.RESULTS) -> dict[str, dict[str, float]]:
    """% de repetições em 1.º lugar no Monte Carlo das propriedades: {material: {método: %}}.

    Em linguagem simples: lê do results/monte_carlo.csv, para um caso, quantas vezes (em %) cada material ficou em
    1.º lugar com cada método."""
    mc = pd.read_csv(results / "monte_carlo.csv")
    t = mc[(mc["caso"] == caso) & (mc["analise"] == "MC propriedades")]
    return {m: g.set_index("metodo")["pct_primeiro"].astype(float).to_dict() for m, g in t.groupby("material")}


def intervalo(material: str, propriedade: str, data: Path = ap.DATA) -> tuple[float, float, str]:
    """Mínimo, máximo e unidade de uma propriedade de um material na base verificada (primeira linha)."""
    m = pd.read_csv(data / "materiais.csv")
    r = m[(m["material"] == material) & (m["propriedade"] == propriedade)].iloc[0]
    return float(r["min"]), float(r["max"]), str(r["unidade"])


FEA_CSV = ap.RESULTS / "fea_haste.csv"
FEA_MAPA = ap.RESULTS / "figures" / "fea" / "sed_corte_coronal_diap.png"
FEA_ORDEM = ["Ti-6Al-4V poroso", "Ti-6Al-4V", "316L", "Co-Cr-Mo"]
FEA_NOMES = {"Ti-6Al-4V poroso": "Liga de titânio porosa\n(17 GPa)", "Ti-6Al-4V": "Liga de titânio\n(Ti-6Al-4V)",
             "316L": "Aço inoxidável\n(316L)", "Co-Cr-Mo": "Cobalto-crómio\n(Co-Cr-Mo)"}


def razoes_fea(csv: Path = FEA_CSV) -> dict[str, tuple[float, float]]:
    """Razão implantado / intacto da energia no osso (zonas de Gruen 1 e 7) na malha mais fina da FEA em Python."""
    d = pd.read_csv(csv)
    d = d[(d["malha_mm"] == d["malha_mm"].min()) & (d["material"].isin(FEA_ORDEM))]
    p = d.pivot(index="material", columns="zona", values="razao")
    return {m: (float(p.loc[m, 1]), float(p.loc[m, 7])) for m in FEA_ORDEM if m in p.index}


def margem_leitura_literal(data: Path = ap.DATA) -> tuple[str, str, float]:
    """(1.º, 2.º, ΔC) do TOPSIS do scaffold na leitura literal do problema crítico (cenário P-foco, só ligação direta).

    Em linguagem simples: é o ranking que o gráfico do diapositivo do scaffold mostra (peso duplicado do suporte
    celular e da degradação); devolve quem fica à frente, quem fica em 2.º e a diferença de C entre os dois, para os
    cartões dizerem a margem da mesma leitura que o gráfico."""
    import animacoes as an
    t = an.leituras_scaffold(data)["literal"].sort_values("TOPSIS", ascending=False)
    return str(t["material"].iloc[0]), str(t["material"].iloc[1]), float(t["TOPSIS"].iloc[0] - t["TOPSIS"].iloc[1])


# Frases únicas das propostas do stent (D1) e do scaffold (D7), decididas pelo André a 11 out 2026 (stent revisto a
# 11 out): a longa no quadro e nas notas; a curta nos títulos e no espectro, a mesma da aplicação.
FRASE_STENT = ("Pt-Cr e L605 praticamente equivalentes: o Pt-Cr é favorecido pelo TOPSIS, pelo VIKOR e pela "
               "radiopacidade; o L605 pelo WASPAS e por ter todas as propriedades em fichas de fornecedor. Sem vencedor "
               "universal.")
FRASE_STENT_CURTA = "Stent: Pt-Cr e L605 praticamente equivalentes"
FRASE_SCAFFOLD_CURTA = "Scaffold: depende da carga (compósito com carga; β-TCP sem carga)"
FRASE_SCAFFOLD = ("Depende da carga: compósito PCL/β-TCP impresso quando o defeito suporta carga (1.º no ranking direto "
                  "em 5 das 6 combinações); β-TCP poroso quando não suporta (mais vezes 1.º na incerteza dos dados).")
COMPOSITO, BTCP = "Compósito PCL/β-TCP (impressão 3D)", "β-TCP poroso"
METODOS = ("TOPSIS", "WASPAS", "VIKOR")


def sem_prefixo(frase: str) -> str:
    """Tira o prefixo "Caso: " de uma frase curta (ex.: "Stent: Pt-Cr e L605..." passa a "Pt-Cr e L605...").

    Em linguagem simples: onde o nome do caso já está escrito ao lado (no espectro), mostra-se o resto da frase,
    exatamente igual."""
    i = frase.find(": ")
    return frase if i < 0 else frase[i + 2:]


def sensibilidade_perfil(perfil: str = "H1", results: Path = ap.RESULTS) -> tuple[str, str, str, str]:
    """(descrição, 1.º do TOPSIS sem perfil, 1.º do TOPSIS com o perfil, 1.º do Borda com o perfil), cenário A.

    Em linguagem simples: no jovem ativo (H1), o fator de peso do perfil põe outro material em 1.º no TOPSIS; o
    diapositivo dos perfis mostra-o como sensibilidade a esse fator, não como recomendação."""
    p = pd.read_csv(results / "perfis_referencia.csv")
    g = p[(p["perfil"] == perfil) & (p["cenario"] == "A")]
    primeiro = lambda col: str(g.loc[g[col] == 1, "material"].iloc[0])  # noqa: E731
    return str(g["descricao"].iloc[0]), primeiro("pos_ref_TOPSIS"), primeiro("pos_TOPSIS"), primeiro("pos_Borda")


def vencedores_ranking(caso: str, results: Path = ap.RESULTS) -> dict[tuple[str, str], str]:
    """1.º lugar do ranking direto por (cenário, método), lido de results/pontuacoes.csv (cenários A e B).

    Em linguagem simples: para cada cenário e cada um dos três métodos, diz que material ficou em 1.º lugar; é daqui
    que vem a contagem "em 5 das 6 combinações" da frase do scaffold."""
    p = pd.read_csv(results / "pontuacoes.csv")
    p = p[p["caso"] == caso]
    return {(c, m): str(g.loc[g[f"pos_{m}"] == 1, "material"].iloc[0])
            for c, g in p.groupby("cenario") for m in METODOS}


def n_primeiro(vencedores: dict[tuple[str, str], str], material: str) -> int:
    """Em quantas combinações (cenário, método) o material fica em 1.º lugar."""
    return sum(v == material for v in vencedores.values())

# Resposta ao enunciado (Caso | Problema | Biomateriais a investigar | Proposta): texto do relatório (Tabela 1 e
# questões 8), com o problema crítico nas palavras do docente.
RESPOSTA = [
    ("Haste femoral", "Stress shielding e fadiga (problema acrescentado)",
     "Aço 316L, Co-Cr-Mo, Ti-6Al-4V ELI, Ti grau 4, Ti-13Nb-13Zr",
     "Ti-6Al-4V ELI, haste não cimentada; porosidade graduada como passo seguinte"),
    ("Par articular", "Desgaste e partículas",
     "CoCrMo/UHMWPE, CoCrMo/HXLPE, ZTA/HXLPE, ZTA/ZTA (metal-metal eliminado)",
     "ZTA/ZTA (indicativo); ZTA/HXLPE com menos risco de fratura e ruído"),
    ("Stent coronário", "Resistência mecânica e biocompatibilidade",
     "316L, Co-Cr L605, MP35N, Pt-Cr; Mg WE43 e PLLA no cenário B",
     FRASE_STENT),
    ("Scaffold ósseo", "Suporte celular e degradação controlada",
     "PCL, HA, β-TCP, vidro 45S5, PCL/β-TCP; PLLA, PLGA e quitosano/HA no cenário B",
     FRASE_SCAFFOLD),
    ("Implante dentário", "Integração com o osso e corrosão",
     "Ti grau 4, Ti-6Al-4V ELI, Ti-Zr, zircónia Y-TZP, PEEK",
     "Ti grau 4 rugoso (indicativo); Ti-Zr em diâmetro reduzido; Y-TZP estética"),
]

LARG_RESPOSTA = (1.45, 2.1, 3.4, 5.18)  # larguras das colunas (pol.): no máximo 3 linhas de 14 pt por célula
# alturas das linhas do quadro da resposta: as do stent e do scaffold levam as frases únicas, com até 5 linhas de 14 pt
ALT_RESPOSTA = dict(y0=1.32, h_linha=0.8, alturas={2: 1.3, 3: 1.2})  # y0: o quadro sobe para caber

# Fabrico, testes e caracterização da proposta de cada caso (relatório, questões 9, 10 e 11 de cada caso).
FABRICO = [
    ("Haste\n(Ti-6Al-4V ELI)", "Forjamento e maquinagem; jateamento e HA por plasma; porosa: fabrico aditivo e HIP",
     "ISO 10993; fadiga ISO 7206-4 e 7206-6; corrosão ASTM F2129",
     "Tração, fadiga, dureza Vickers, SEM/EDS, XRD (fases α e β), XPS"),
    ("Par articular\n(ZTA/ZTA)", "Prensagem do pó, sinterização, HIP, retificação e polimento",
     "Cabeça ISO 7206-10; desgaste em simulador (ISO 14242-1 e 14242-2)",
     "Densidade, dureza, tenacidade, XRD (fase monoclínica), rugosidade"),
    ("Stent\n(Pt-Cr, L605)", "Tubo trefilado, corte a laser, recozimento, eletropolimento, fármaco",
     "ISO 25539-2; fadiga pulsátil ASTM F2477; recuo ASTM F2079",
     "Tração do tubo, SEM dos struts, ICP-MS (níquel), radiopacidade"),
    ("Scaffold\n(PCL/β-TCP)", "Impressão 3D por deposição de material fundido, com poros interligados",
     "ASTM F2150 e F2450; degradação ISO 13781; defeito ósseo em animal",
     "Micro-CT (porosidade), SEM, compressão, GPC e DSC"),
    ("Dentário\n(Ti grau 4)", "Maquinagem, jateamento e ataque ácido; zircónia: pó sinterizado",
     "Fadiga ISO 14801; ISO 10993; zircónia ISO 13356",
     "Rugosidade (Sa), SEM e XPS, ângulo de contacto, % BIC in vivo"),
]


def gravar_diap(fig, destino: Path) -> Path:
    """Grava a figura com o tamanho exato que tem no diapositivo (sem recorte), em PNG a 300 ppp e fundo branco.

    Em linguagem simples: o estilo do projeto recorta as margens ao gravar, o que muda o tamanho da imagem e, ao
    encaixá-la na caixa do diapositivo, encolhe a letra; aqui a imagem fica com as polegadas da caixa, por isso 14 pt
    no gráfico são 14 pt no diapositivo."""
    import matplotlib.pyplot as plt
    destino.parent.mkdir(parents=True, exist_ok=True)
    with plt.rc_context({"savefig.bbox": "standard"}):  # o estilo do projeto usa "tight", que recorta
        fig.savefig(destino, dpi=300, transparent=False, facecolor="white")
    plt.close(fig)
    return destino


def grafico_fea(razoes: dict[str, tuple[float, float]], destino: Path, w: float = 5.25, h: float = 4.3) -> Path:
    """Gráfico das razões em Gruen 1 e 7, desenhado ao tamanho do diapositivo (letra mínima de 14 pt)."""
    import estilo_graficos as eg
    import matplotlib.pyplot as plt
    import numpy as np
    eg.aplicar("clara")
    fig, ax = plt.subplots(figsize=(w, h), dpi=300)
    mats = [m for m in FEA_ORDEM if m in razoes]
    y = np.arange(len(mats))[::-1]
    g1 = [razoes[m][0] for m in mats]
    g7 = [razoes[m][1] for m in mats]
    ax.barh(y + 0.19, g1, 0.36, color="#0072B2", label="Gruen 1")
    ax.barh(y - 0.19, g7, 0.36, color="#E69F00", label="Gruen 7 (colo)")
    ax.axvline(1, color="#55595e", lw=1, ls="--")
    ax.text(1.02, len(mats) - 0.45, "osso sem\nprótese", fontsize=14, color="#55595e", va="top")
    ax.set_yticks(y, [FEA_NOMES[m] for m in mats], fontsize=14)
    ax.set_xlim(0, 1.45)
    ax.set_xticks([0, 0.5, 1], ["0", "0,5", "1"], fontsize=14)
    ax.set_xlabel("Energia no osso" + chr(10) + "(1 = osso sem prótese)", fontsize=14)
    ax.legend(fontsize=14, loc="lower left", bbox_to_anchor=(-0.02, 1.0), ncol=2, frameon=False, handlelength=1.0,
              columnspacing=0.8)
    ax.grid(axis="y", visible=False)
    fig.tight_layout()
    return gravar_diap(fig, destino)


CFD_JSON = ap.RESULTS / "cfd_pulsatil.json"
CFD_MAPA = ap.RESULTS / "figures" / "cfd" / "tawss_mapas_diap.png"
CFD_ORDEM = ["Pt-Cr", "316L"]
CFD_NOMES = {"Pt-Cr": "Strut fino\n(79 µm)", "316L": "Strut grosso\n(135 µm)"}


def dados_cfd(js: Path = CFD_JSON) -> dict[str, dict]:
    """Resultados do fluxo pulsátil (results/cfd_pulsatil.json, gerado por scripts/cfd_stent.py resumo)."""
    import json
    return json.loads(js.read_text(encoding="utf-8"))


def grafico_cfd(dados: dict[str, dict], destino: Path, w: float = 5.25, h: float = 4.3) -> Path:
    """Barras da área da parede com corte baixo (TAWSS abaixo de 0,4 Pa), ao tamanho do diapositivo."""
    import estilo_graficos as eg
    import matplotlib.pyplot as plt
    import numpy as np
    eg.aplicar("clara")
    fig, ax = plt.subplots(figsize=(w, h), dpi=300)
    mats = [m for m in CFD_ORDEM if m in dados]
    v = [100 * dados[m]["fracao_baixa"] for m in mats]
    x = np.arange(len(mats))
    ax.bar(x, v, 0.55, color="#E67100")
    for xi, vi in zip(x, v):
        ax.text(xi, vi + 1.5, f"{vi:.0f} %", ha="center", va="bottom", fontsize=16, fontweight="bold")
    ax.set_xticks(x, [CFD_NOMES[m].replace("(79 µm)", f"({dados[m]['espessura_um']:.0f} µm)").replace(
        "(135 µm)", f"({dados[m]['espessura_um']:.0f} µm)") for m in mats], fontsize=14)
    ax.set_ylim(0, max(v) * 1.25 + 5)
    ax.set_yticks([])
    for lado in ("top", "right", "left"):
        ax.spines[lado].set_visible(False)
    ax.set_title("Parede com fluxo lento (abaixo de\n0,4 Pa), sangue newtoniano", fontsize=14, loc="left")
    ax.grid(visible=False)
    fig.tight_layout()
    return gravar_diap(fig, destino)


# Duas ferramentas independentes, o mesmo resultado (results/fea_haste_mapdl.md e results/cfd_fluent.md, ramo
# extensoes, 11 out 2026): haste no SOLIDWORKS (malha de 3,5 mm) e no Ansys Mechanical (MAPDL, 3,0 mm), no mesmo
# fémur simplificado; stent no OpenFOAM e no Ansys Fluent, com a mesma malha (fluxo estacionário ao caudal médio).
SW_JSON = ap.RESULTS / "fea_haste_sw.json"
MAPDL_JSON = ap.RESULTS / "fea_haste_mapdl.json"
FLUENT_JSON = ap.RESULTS / "cfd_fluent.json"
NOMES_FEA_DUAS = {"Ti-6Al-4V poroso": "Ti-6Al-4V poroso", "Ti-6Al-4V ELI": "Ti-6Al-4V ELI",
                  "Aço inox 316L": "316L", "Co-Cr-Mo forjado": "Co-Cr-Mo"}


def razoes_duas_ferramentas(sw: Path = SW_JSON, mapdl: Path = MAPDL_JSON,
                            malha_sw: str = "3.5") -> dict[str, dict[str, tuple[float, float]]]:
    """Razões implantado / intacto da energia no osso em Gruen 1 e 7, no SOLIDWORKS e no Ansys (MAPDL).

    Em linguagem simples: para cada material, divide a energia de deformação no osso das zonas 1 e 7 com a haste
    pela do fémur sem prótese, nas duas ferramentas; devolve {material: {"SOLIDWORKS": (g1, g7), "Ansys": (g1, g7)}}.
    """
    import json
    s = json.loads(sw.read_text(encoding="utf-8"))["casos"]
    m = json.loads(mapdl.read_text(encoding="utf-8"))

    def razao(caso: dict, intacto: dict, z: str) -> float:
        return caso["zonas"][z]["energia_J"] / intacto["zonas"][z]["energia_J"]

    out = {}
    for mat in NOMES_FEA_DUAS:
        if mat in s and mat in m:
            out[mat] = {"SOLIDWORKS": tuple(razao(s[mat][malha_sw], s["Intacto"][malha_sw], z) for z in "17"),
                        "Ansys": tuple(razao(m[mat], m["Intacto"], z) for z in "17")}
    return out


def dif_max_razoes(r: dict[str, dict[str, tuple[float, float]]]) -> float:
    """Maior diferença relativa (%) entre o Ansys e o SOLIDWORKS nas razões de Gruen 1 e 7."""
    return max(abs(a / b - 1) * 100 for v in r.values() for a, b in zip(v["Ansys"], v["SOLIDWORKS"]))


def areas_duas_ferramentas(js: Path = FLUENT_JSON) -> dict[str, dict[str, float]]:
    """% da parede com tensão de corte abaixo de 0,4 Pa no OpenFOAM e no Fluent: {liga: {ferramenta: %}}."""
    import json
    d = json.loads(js.read_text(encoding="utf-8"))
    return {liga: {"OpenFOAM": 100 * v["of_fracao_baixa"], "Ansys Fluent": 100 * v["fracao_baixa"]}
            for liga, v in d.items()}


def grafico_duas_ferramentas(haste: dict, stent: dict, destino: Path, w: float = 12.13, h: float = 3.6) -> Path:
    """Barras emparelhadas (uma ferramenta contra a outra) da haste e do stent, ao tamanho do diapositivo (14 pt).

    Em linguagem simples: à esquerda, a energia no osso junto ao colo (Gruen 7) com cada material, no SOLIDWORKS
    e no Ansys Mechanical; à direita, a parede do stent com fluxo lento no OpenFOAM e no Ansys Fluent. Barras do
    mesmo par lado a lado e quase iguais: as duas ferramentas dão a mesma resposta."""
    import estilo_graficos as eg
    import matplotlib.pyplot as plt
    import numpy as np
    from figuras import AZUL, LARANJA
    eg.aplicar("clara")
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(w, h), dpi=300, gridspec_kw={"width_ratios": [1.45, 1]})
    mats = [m for m in NOMES_FEA_DUAS if m in haste]
    y = np.arange(len(mats))[::-1]
    for ferr, cor, dy in (("SOLIDWORKS", AZUL, 0.19), ("Ansys", LARANJA, -0.19)):
        v = [haste[m][ferr][1] for m in mats]
        a1.barh(y + dy, v, 0.36, color=cor, label="Ansys Mechanical" if ferr == "Ansys" else ferr)
        for yi, vi in zip(y, v):
            a1.text(vi + 0.03, yi + dy, eg.numero(vi, 2), va="center", fontsize=14)
    a1.set_yticks(y, [NOMES_FEA_DUAS[m] for m in mats], fontsize=14)
    a1.set_xlim(0, 1.9)
    a1.set_xticks([0, 0.5, 1, 1.5], ["0", "0,5", "1", "1,5"], fontsize=14)
    a1.set_xlabel("Haste: energia no osso junto ao colo (Gruen 7; 1 = osso sem prótese)", fontsize=14)
    a1.legend(fontsize=14, loc="lower right", frameon=False, handlelength=1.0)
    a1.grid(axis="y", visible=False)
    ligas = [x for x in ("Pt-Cr", "316L") if x in stent]
    x = np.arange(len(ligas))
    for ferr, cor, dx in (("OpenFOAM", AZUL, -0.2), ("Ansys Fluent", LARANJA, 0.2)):
        v = [stent[li][ferr] for li in ligas]
        a2.bar(x + dx, v, 0.38, color=cor, label=ferr)
        for xi, vi in zip(x, v):
            a2.text(xi + dx, vi + 2, nb(f"{eg.numero(vi, 1)} %"), ha="center", va="bottom", fontsize=14)
    a2.set_xticks(x, [{"Pt-Cr": "Pt-Cr (79 µm)", "316L": "316L (135 µm)"}[li] for li in ligas], fontsize=14)
    a2.set_ylim(0, 118)
    a2.set_yticks([0, 50, 100], ["0 %", "50 %", "100 %"], fontsize=14)
    a2.set_xlabel("Stent: parede abaixo de 0,4 Pa", fontsize=14)
    a2.legend(fontsize=14, loc="upper left", frameon=False, handlelength=1.0, ncol=2)
    a2.grid(axis="x", visible=False)
    fig.tight_layout(w_pad=2.0)
    return gravar_diap(fig, destino)


def primeiro_existente(candidatos: list[Path]) -> Path | None:
    """O primeiro ficheiro da lista que existe (None se nenhum existe): para o logótipo e o vídeo da demo."""
    return next((p for p in candidatos if p.exists()), None)


def capa_video(mp4: Path, destino: Path) -> Path:
    """Imagem da última fotografia do vídeo (capa do vídeo no diapositivo e no PDF), com o ffmpeg do imageio."""
    import subprocess

    import imageio_ffmpeg
    _, seg = imageio_ffmpeg.count_frames_and_secs(str(mp4))
    destino.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-y", "-loglevel", "error", "-ss", f"{max(0.0, seg - 0.2):.2f}",
                    "-i", str(mp4), "-frames:v", "1", str(destino)], check=True)
    return destino


def palavras_notas(prs, n: int = N_PRINCIPAIS) -> int:
    """Número de palavras das notas do orador dos n primeiros diapositivos (os de reserva não contam)."""
    return sum(len(s.notes_slide.notes_text_frame.text.split()) for s in list(prs.slides)[:n] if s.has_notes_slide)


def minutos_fala(palavras: int, segundos_video: float = 0.0, ppm: int = PALAVRAS_POR_MINUTO) -> float:
    """Tempo estimado (min): as palavras ao ritmo do projeto mais os vídeos em que não se fala."""
    return palavras / ppm + segundos_video / 60


def grafico_funil(tab: pd.DataFrame, destino: Path, w: float = 7.9, h: float = 4.9) -> Path:
    """Funil de triagem ao tamanho do diapositivo: números a 18 pt e legendas a 14 pt.

    Em linguagem simples: as mesmas etapas e larguras de scripts/graficos.py (largura pelo número de algarismos), mas
    com letra legível no diapositivo; na base de dados, a legenda diz que cada valor tem fonte e estado (nem todos
    estão verificados)."""
    import estilo_graficos as eg
    import graficos as gr
    import matplotlib.pyplot as plt
    import numpy as np
    t = eg.aplicar("clara")
    d = gr.dados_funil(tab)
    base = d["etapa"] == "Base verificada (materiais por componente)"
    d.loc[base, "texto"] = "escolhidos à parte: normas ASTM e uso clínico"
    fig = plt.figure(figsize=(w, h), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1])
    cor_ini, cor_fim = np.array([0.62, 0.66, 0.70]), np.array([0.0, 0.447, 0.698])
    k = len(d)
    if base.any():  # a passagem para a base não é um filtro automático: o funil parte-se aqui
        yq = k - 1 - int(np.flatnonzero(base.to_numpy())[0]) + 0.5
        ax.plot([-0.53, 0.53], [yq, yq], color=t["texto2"], lw=1.6, ls=(0, (4, 3)))
    for i, r in d.iterrows():
        y = k - 1 - i
        cor = cor_ini + (cor_fim - cor_ini) * (i / (k - 1))
        largura = 0.08 + 0.92 * r["largura"]
        ax.barh(y, largura, left=-largura / 2, height=0.7, color=cor, lw=0)
        ax.text(0.56, y + 0.13, eg.numero(r["n"], 0), ha="left", va="center", fontsize=18, fontweight="bold",
                color=t["texto"])
        ax.text(0.56, y - 0.22, r["texto"], ha="left", va="center", fontsize=14, color=t["texto2"])
    ax.set_xlim(-0.55, 2.1)
    ax.set_ylim(-0.6, k - 0.4)
    ax.axis("off")
    return gravar_diap(fig, destino)


def grafico_mc_stent(pct: dict[str, dict[str, float]], destino: Path, w: float = 5.55, h: float = 4.7) -> Path:
    """% de repetições em 1.º lugar no Monte Carlo do stent (três métodos), ao tamanho do diapositivo (14 pt).

    Em linguagem simples: uma linha por material, do método em que ganha menos ao método em que ganha mais, com uma
    marca por método e a faixa escrita ao lado (por exemplo, "56 a 70 %")."""
    import estilo_graficos as eg
    import matplotlib.pyplot as plt
    from figuras import COR_MATERIAL
    t = eg.aplicar("clara")
    metodos, marcas = ("TOPSIS", "WASPAS", "VIKOR"), {"TOPSIS": "o", "WASPAS": "D", "VIKOR": "s"}
    mats = sorted(pct, key=lambda m: pct[m].get("TOPSIS", 0))
    fig, ax = plt.subplots(figsize=(w, h), dpi=300)
    for i, m in enumerate(mats):
        v = [pct[m][k] for k in metodos]
        lo, hi = min(v), max(v)
        cor = COR_MATERIAL["Stent"].get(m, t["texto2"])
        ax.plot([lo, hi], [i, i], color=cor, lw=9, alpha=0.25, solid_capstyle="round")
        for k in metodos:
            ax.scatter(pct[m][k], i, marker=marcas[k], s=90, color=cor, zorder=3, edgecolor="white", linewidth=1.0)
        faixa = f"{eg.numero(lo, 0)} a {eg.numero(hi, 0)} %" if round(hi) != round(lo) else f"{eg.numero(hi, 0)} %"
        ax.annotate(nb(faixa), xy=(hi, i), xytext=(12, 0), textcoords="offset points", va="center", fontsize=14,
                    fontweight="bold" if i == len(mats) - 1 else "normal", color=t["texto"])
    ax.set_yticks(range(len(mats)), [ns.com_codigo(m).replace(" (", "\n(") for m in mats], fontsize=14)
    ax.set_xlim(-3, 100)
    ax.set_xticks([0, 50, 100], ["0 %", "50 %", "100 %"], fontsize=14)
    ax.set_xlabel("Repetições em 1.º lugar", fontsize=14)
    ax.grid(axis="y", visible=False)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0)
    for k in metodos:
        ax.scatter([], [], marker=marcas[k], color=t["texto2"], s=60, label=k)
    ax.legend(loc="lower right", bbox_to_anchor=(1.0, 1.0), ncol=3, handletextpad=0.1, columnspacing=0.5,
              fontsize=14, frameon=False)
    fig.tight_layout()
    return gravar_diap(fig, destino)


def grafico_scaffold(vencedores: dict[tuple[str, str], str], pct: dict[str, dict[str, float]], destino: Path,
                     w: float = 8.6, h: float = 4.84) -> Path:
    """Scaffold ao tamanho do diapositivo (14 pt): à esquerda, o 1.º lugar do ranking direto por cenário e método;
    à direita, as repetições em 1.º lugar no Monte Carlo dos dados (cenário A), por método.

    Em linguagem simples: a grelha mostra quem ganha quando se usam os valores típicos (o compósito em quase todas as
    casas); as barras mostram quem ganha mais vezes quando os dados variam dentro dos intervalos (o β-TCP)."""
    import estilo_graficos as eg
    import matplotlib.pyplot as plt
    from figuras import COR_MATERIAL
    t = eg.aplicar("clara")
    cores = COR_MATERIAL["Scaffold"]
    rotulo = {COMPOSITO: "Compósito", BTCP: "β-TCP"}
    fig = plt.figure(figsize=(w, h), dpi=300)
    ax = fig.add_axes([0.16, 0.17, 0.30, 0.62])
    cenarios = sorted({c for c, _ in vencedores})
    for j, c in enumerate(cenarios):  # uma coluna por cenário, uma linha por método
        for i, m in enumerate(METODOS):
            v = vencedores[(c, m)]
            ax.add_patch(plt.Rectangle((j + 0.04, i + 0.05), 0.92, 0.9, color=cores.get(v, t["texto2"]), lw=0))
            ax.text(j + 0.5, i + 0.5, rotulo.get(v, curto(v)), ha="center", va="center", fontsize=14,
                    fontweight="bold", color="white" if v == BTCP else t["texto"])
    ax.set_xlim(0, len(cenarios))
    ax.set_ylim(len(METODOS), 0)
    ax.set_xticks([j + 0.5 for j in range(len(cenarios))], [f"Cenário {c}" for c in cenarios], fontsize=14)
    ax.set_yticks([i + 0.5 for i in range(len(METODOS))], METODOS, fontsize=14)
    ax.tick_params(length=0)
    ax.xaxis.tick_top()
    for s in ax.spines.values():
        s.set_visible(False)
    ax.grid(False)
    fig.text(0.02, 0.95, "1.º lugar no ranking direto", fontsize=14, fontweight="bold", color=t["texto"])
    ax.text(0, len(METODOS) + 0.3, nb(f"compósito em {n_primeiro(vencedores, COMPOSITO)} de {len(vencedores)}"),
            ha="left", va="top", fontsize=14, color=t["texto2"])
    bx = fig.add_axes([0.62, 0.2, 0.33, 0.55])
    mats = [m for m in (BTCP, COMPOSITO) if m in pct]
    alt = 0.36
    for i, m in enumerate(mats):
        for j, k in enumerate(METODOS):
            y = j + (i - 0.5) * alt
            bx.barh(y, pct[m][k], height=alt * 0.92, color=cores[m], lw=0)
            bx.text(pct[m][k] + 2, y, nb(f"{eg.numero(pct[m][k], 0)} %"), va="center", fontsize=14,
                    color=t["texto"])
    bx.set_yticks(range(len(METODOS)), METODOS, fontsize=14)
    bx.set_ylim(len(METODOS) - 0.4, -0.6)
    bx.set_xlim(0, 100)
    bx.set_xticks([0, 50, 100], ["0 %", "50 %", "100 %"], fontsize=14)
    bx.set_xlabel("repetições em 1.º lugar", fontsize=14)
    bx.tick_params(axis="y", length=0)
    bx.grid(axis="y", visible=False)
    bx.spines["left"].set_visible(False)
    fig.text(0.52, 0.95, "Monte Carlo dos dados (cenário A)", fontsize=14, fontweight="bold", color=t["texto"])
    for i, m in enumerate(mats):
        fig.text(0.62 + i * 0.2, 0.84, "■ " + rotulo[m], fontsize=14, color=cores[m], fontweight="bold")
    return gravar_diap(fig, destino)


def grafico_gibson(destino: Path, w: float = 5.55, h: float = 4.7, data: Path = ap.DATA) -> Path:
    """Gibson e Ashby, E = E maciço × (1 − p)², ao tamanho do diapositivo (14 pt), com a faixa do osso cortical.

    Em linguagem simples: a curva mostra quanto a rigidez da liga de titânio cai com a porosidade; o ponto marca a
    porosidade em que chega ao alvo do osso. Os rótulos ficam em zonas livres para não se taparem."""
    import animacoes as an
    import estilo_graficos as eg
    import matplotlib.pyplot as plt
    import numpy as np
    from figuras import COR_MATERIAL
    t = eg.aplicar("clara")
    e_ti, alvo = an.modulo_e_alvo_haste(data)
    tec = pd.read_csv(data / "tecido.csv")
    cort = tec[(tec["tecido"] == "Osso cortical") & (tec["propriedade"] == "Módulo de Young")].iloc[0]
    cor = COR_MATERIAL["Haste"]["Ti-6Al-4V ELI"]
    p = np.linspace(0, 0.85, 200)
    ps = an.porosidade_para(alvo, e_ti)
    fig, ax = plt.subplots(figsize=(w, h), dpi=300)
    ax.axhspan(cort["min"], cort["max"], color="#E69F00", alpha=0.18, lw=0)
    ax.text(1, cort["max"] + 2, nb(f"osso cortical ({eg.numero(cort['min'], 0)} a {eg.numero(cort['max'], 0)} GPa)"),
            ha="left", va="bottom", fontsize=14, color=t["texto2"])
    ax.plot(p * 100, an.gibson_ashby(p, e_ti), color=cor, lw=2.6)
    ax.plot([ps * 100, ps * 100], [0, alvo], color=t["texto2"], lw=1.0, ls=(0, (3, 3)))
    ax.scatter([ps * 100], [alvo], s=90, color=cor, zorder=4, edgecolor="white", linewidth=1.2)
    ax.annotate(nb(f"{eg.numero(alvo, 0)} GPa") + "\n" + nb(f"com p = {eg.numero(ps * 100, 0)} %"),
                xy=(ps * 100, alvo), xytext=(8, 22), textcoords="offset points", ha="left", va="bottom", fontsize=14,
                fontweight="bold", color=t["texto"])
    ax.annotate(nb(f"Ti-6Al-4V maciço: {eg.numero(e_ti, 1)} GPa"), xy=(0, e_ti), xytext=(14, 4),
                textcoords="offset points", ha="left", va="bottom", fontsize=14, color=t["texto"])
    ax.set_xlim(0, 85)
    ax.set_ylim(0, e_ti * 1.22)
    ax.set_xticks([0, 20, 40, 60, 80], [f"{v} %" for v in (0, 20, 40, 60, 80)], fontsize=14)
    ax.set_yticks([0, 50, 100], ["0", "50", "100"], fontsize=14)
    ax.set_xlabel("Porosidade p", fontsize=14)
    ax.set_ylabel("Módulo E (GPa)", fontsize=14)
    fig.tight_layout()
    return gravar_diap(fig, destino)


def nb(s: str) -> str:
    """Troca os espaços por espaços que não partem a linha (para "10 000" ou "9 999 | 10 000" ficarem juntos)."""
    return str(s).replace(" ", "\u00a0").replace("\u2009", "\u00a0")


def mil(v) -> str:
    return nb(ap.milhares(v))


def n_regras(data: Path = ap.DATA) -> int:
    return len(pd.read_csv(data / "perfis_doente.csv"))


def imagem_reduzida(origem: Path, largura: int, destino_pasta: Path = CACHE) -> Path:
    """Cópia da imagem com a largura dada (os renders 4K pesariam dezenas de MB na apresentação)."""
    destino_pasta.mkdir(parents=True, exist_ok=True)
    d = destino_pasta / f"{origem.stem}_{largura}.png"
    if not d.exists() or d.stat().st_mtime < origem.stat().st_mtime:
        im = Image.open(origem)
        if im.width > largura:
            im = im.resize((largura, round(im.height * largura / im.width)), Image.LANCZOS)
        im.save(d, optimize=True)
    return d


def recortar_transparente(origem: Path, destino_pasta: Path = CACHE, margem: float = 0.06,
                          esbater: float = 0.10) -> Path:
    """Recorta a imagem transparente ao objeto e esbate a sombra junto às margens do recorte.

    Em linguagem simples: a caixa do recorte segue só o objeto (píxeis quase opacos), com uma margem; a sombra de
    contacto, que é meio transparente e se estende para lá do objeto, vai ficando transparente perto das margens,
    para não aparecer cortada a direito sobre fundos escuros."""
    import numpy as np
    destino_pasta.mkdir(parents=True, exist_ok=True)
    d = destino_pasta / f"{origem.stem}_recorte.png"
    if not d.exists() or d.stat().st_mtime < origem.stat().st_mtime:
        im = Image.open(origem).convert("RGBA")
        objeto = im.getchannel("A").point(lambda a: 255 if a > 230 else 0)
        x0, y0, x1, y1 = objeto.getbbox() or (0, 0, im.width, im.height)
        m = round(margem * max(x1 - x0, y1 - y0))
        im = im.crop((max(0, x0 - m), max(0, y0 - m), min(im.width, x1 + m), min(im.height, y1 + m)))
        a = np.asarray(im, dtype=np.float32).copy()
        h, w = a.shape[:2]
        e = max(1.0, esbater * max(w, h))
        dx = np.minimum(np.arange(w), np.arange(w)[::-1]) / e
        dy = np.minimum(np.arange(h), np.arange(h)[::-1]) / e
        rampa = np.clip(np.minimum.outer(dy, dx), 0, 1)
        sombra = a[..., 3] < 230  # só a sombra se esbate; o objeto fica intacto
        a[..., 3] = np.where(sombra, a[..., 3] * rampa, a[..., 3])
        Image.fromarray(a.astype(np.uint8), "RGBA").save(d)
    return d


def logotipo_branco(origem: Path, destino_pasta: Path = CACHE) -> Path:
    """Versão branca de um logótipo com transparência (para a capa escura)."""
    destino_pasta.mkdir(parents=True, exist_ok=True)
    d = destino_pasta / f"{origem.stem}_branco.png"
    im = Image.open(origem).convert("RGBA")
    branco = Image.new("RGBA", im.size, (255, 255, 255, 0))
    branco.putalpha(im.getchannel("A"))
    branco.save(d)
    return d


# ---------------------------------------------------------------- desenho

def caixa(slide, x, y, w, h, linhas, tamanho=18, cor=TEXTO, negrito=False, alinha=PP_ALIGN.LEFT,
          ancora=MSO_ANCHOR.TOP, nome=None, espaco=1.0):
    """Caixa de texto na letra do estilo; `linhas` é uma str ou uma lista de str ou de (str, tamanho, cor, negrito)."""
    return ap.texto(slide, x, y, w, h, linhas, tamanho, cor, negrito, LETRA, alinha, ancora, nome, espaco)


def cartao(slide, x, y, w, h, cor=CARTAO, nome=None):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.adjustments[0] = min(0.12, 0.16 / min(w, h))
    c.fill.solid()
    c.fill.fore_color.rgb = cor
    c.line.fill.background()
    c.shadow.inherit = False
    if nome:
        c.name = nome
    return c


def quadro(slide, cabecalho: tuple[str, ...], linhas: list[tuple[str, ...]], larguras: tuple[float, ...],
           x0: float = 0.6, y0: float = 1.75, h_cab: float = 0.45, h_linha: float = 0.95,
           alturas: dict[int, float] | None = None) -> list:
    """Quadro de texto (cabeçalho e uma linha por caso) em cartões, com letra de 14 pt (16 pt na 1.ª coluna).

    Em linguagem simples: desenha uma tabela simples, com o cabeçalho a azul e cada linha num cartão claro; devolve
    as caixas do cabeçalho (as que entram com animação), para o quadro aparecer já completo. `alturas` dá a altura
    de linhas com mais texto ({índice: polegadas}); as outras ficam com h_linha."""
    formas, x = [], x0
    for cab, w in zip(cabecalho, larguras):
        formas.append(caixa(slide, x + 0.12, y0, w - 0.2, h_cab, cab, 14, AZUL_TEXTO, True, ancora=MSO_ANCHOR.BOTTOM))
        x += w
    y = y0 + h_cab + 0.06
    for i, linha in enumerate(linhas):
        h_linha_i = (alturas or {}).get(i, h_linha)
        cartao(slide, x0, y, sum(larguras), h_linha_i)
        x = x0
        for j, (txt, w) in enumerate(zip(linha, larguras)):
            partes = [(t, 16 if j == 0 else 14, TEXTO, j == 0) for t in txt.split("\n")]  # \n: nova linha
            caixa(slide, x + 0.12, y + 0.03, w - 0.2, h_linha_i - 0.06, partes, ancora=MSO_ANCHOR.MIDDLE)
            x += w
        y += h_linha_i + 0.05
    return formas


def numero_grande(slide, x, y, w, valor, legenda, cor=AZUL, tam=40, nome=None):
    return caixa(slide, x, y, w, 1.4, [(valor, tam, cor, True), (legenda, 15, CINZA, False)], nome=nome)


def titulo(slide, frase: str, kicker: str | None = None, n: int | None = None, escuro=False) -> list:
    """Título-afirmação no topo (posição fixa), com o kicker por cima e o rodapé; devolve as formas a animar."""
    formas = []
    if kicker:
        formas.append(caixa(slide, 0.6, 0.38, 10.0, 0.4, kicker.upper(), 14, SUB_NOITE if escuro else AZUL_TEXTO,
                            True, nome="!!kicker"))
    formas.append(caixa(slide, 0.6, 0.72, 12.1, 0.9, frase, 30, BRANCO if escuro else TEXTO, True, nome="!!titulo",
                        ancora=MSO_ANCHOR.TOP))
    if n is not None:
        caixa(slide, 0.6, 6.98, 10.0, 0.32, f"{AUTOR}  ·  Qual é o biomaterial ideal?", 14, CINZA, nome="!!rodape")
        caixa(slide, 12.0, 6.98, 0.73, 0.32, str(n), 14, CINZA, alinha=PP_ALIGN.RIGHT, nome="!!numero")
    return formas


def figura(slide, caminho: Path, x, y, w=None, h=None, nome=None):
    f = slide.shapes.add_picture(str(caminho), Inches(x), Inches(y), Inches(w) if w else None,
                                 Inches(h) if h else None)
    if nome:
        f.name = nome
    return f


def figura_caixa(slide, caminho: Path, x, y, w, h, nome=None):
    """Imagem dentro da caixa (x, y, w, h), centrada e sem deformar."""
    im = Image.open(caminho)
    r = im.width / im.height
    if w / h > r:
        ww, hh = h * r, h
    else:
        ww, hh = w, w / r
    return figura(slide, caminho, x + (w - ww) / 2, y + (h - hh) / 2, ww, hh, nome)


def video(slide, nome: str, x, y, w, nome_forma=None):
    """Vídeo 16:9 com a última imagem como capa; devolve (forma, duração em ms)."""
    import imageio_ffmpeg
    mp4 = ap.MEDIA / f"{nome}.mp4"
    _, seg = imageio_ffmpeg.count_frames_and_secs(str(mp4))
    capa = ap.MEDIA / f"{nome}_final.png"
    f = slide.shapes.add_movie(str(mp4), Inches(x), Inches(y), Inches(w), Inches(w * 9 / 16),
                               poster_frame_image=str(capa), mime_type="video/mp4")
    if nome_forma:
        f.name = nome_forma
    return f, int(seg * 1000)


def video_simulacao(slide, nome: str, legenda: str):
    """Vídeo de simulação de reserva (6,9 pol. de largura, à esquerda) com a legenda por baixo.

    Devolve (formas, vídeos); se o ficheiro não existir, devolve duas listas vazias. O vídeo corre ao clicar."""
    if not (ap.MEDIA / f"{nome}.mp4").exists():
        return [], []
    v, ms = video(slide, nome, 0.5, 1.75, 6.9, "!!video_simulacao")
    leg = caixa(slide, 0.5, 5.68, 6.9, 0.48, legenda, 10, CINZA, nome="!!legenda_video")
    return [v, leg], [(v, ms, False)]


LEGENDA_FEA = ("Corte coronal, carga da marcha, interface colada. Deformação ampliada ×5. "
               "Verificação numérica, não previsão clínica.")
LEGENDA_CFD = ("TAWSS: média ao longo do ciclo cardíaco. Linhas de corrente: instante final do ciclo (t = 2,12 s). "
               "O cursor indica a fase do ciclo; o campo não varia no vídeo.")


def fundo(slide, cor):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = cor


# ---------------------------------------------------------------- diapositivos

def construir(numeros: ap.Numeros | None = None, destino: Path = DESTINO, renders: Path | None = None,
              demo: bool = True) -> Path:
    """Gera o pptx; com demo=False, o diapositivo da aplicação leva o marcador mesmo que o vídeo da demo exista."""
    n = numeros or ap.ler_numeros()
    R = renders or pasta_renders()
    fu = funil()
    rend = {c: imagem_reduzida(recortar_transparente(R / f"{c}_transparente.png"), 1400) for c, _, _ in DISPOSITIVOS}
    rend["stent"] = imagem_reduzida(recortar_transparente(R / "stent_transparente.png"), 1400)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    vazio = prs.slide_layouts[6]
    plano: list[tuple] = []  # (slide, entradas, vídeos, esquema)

    def novo(escuro=False):
        s = prs.slides.add_slide(vazio)
        fundo(s, NOITE if escuro else BRANCO)
        return s

    logo = primeiro_existente(LOGO_PROJETO)  # logótipo do projeto (capa e conclusão), quando existir

    # 1. Capa: vídeo de abertura em ecrã inteiro (15 s, automático); o título está no próprio vídeo e a última
    # imagem é a capa; o resto entra quando o vídeo acaba.
    s = novo(True)
    v_ab, d_ab = video(s, ABERTURA, 0, 0, 13.333, "!!video_abertura")
    lu = figura(s, logotipo_branco(ap.LOGOS / "lusofona.png"), 0.6, 0.5, h=0.75)
    fe = figura(s, ap.LOGOS / "fe.png", 1.85, 0.5, h=0.75)
    # cabeçalho em duas linhas, para não tocar no osso da bacia, que no vídeo sobe até ao topo à direita
    cab1, cab2 = CABECALHO.split(" · ", 1)
    k = caixa(s, 3.05, 0.5, 5.9, 0.75, [(cab1, 14, SUB_NOITE, False), (cab2, 14, SUB_NOITE, False)],
              ancora=MSO_ANCHOR.MIDDLE)
    entradas_capa = [lu, fe, k]
    if logo:
        entradas_capa.append(figura(s, logo, 11.95, 0.45, h=0.85, nome="!!logo_projeto"))
    sub = caixa(s, 0.6, 4.75, 7.6, 1.0, "Como escolher o material de cinco implantes com dados verificados e "
                "testes de robustez", 20, SUB_NOITE)
    aut = caixa(s, 0.6, 5.85, 7.9, 0.9, [(AUTOR, 18, BRANCO, True),
                                         (f"Orientador: {ORIENTADOR}, docente de Próteses e Órgãos Artificiais", 15,
                                          SUB_NOITE, False)])
    ap.notas(s, "(Deixar correr o vídeo de abertura, 15 segundos.) Bom dia. Quando um hospital escolhe uma prótese, "
                "escolhe também o material que vai viver dentro do corpo durante anos. A resposta honesta à pergunta "
                "é que depende, e o objetivo foi medir de quê. Comparei materiais para cinco implantes, com dados "
                "verificados e três métodos de decisão, e testei se cada escolha resiste quando os dados mudam. O "
                "trabalho foi orientado pelo professor Pedro Sampaio.")
    plano.append((s, entradas_capa + [sub, aut], [(v_ab, d_ab, True)], "abertura"))

    # 2. O problema
    s = novo()
    f = titulo(s, "Que material pôr dentro do corpo?", "O problema", 2)
    cols = []
    cartao(s, 0.6, 1.85, 12.13, 3.75, NOITE)
    for i, (c, nome, desafio) in enumerate(DISPOSITIVOS):
        x = 0.6 + i * 2.426
        figura_caixa(s, rend[c], x + 0.05, 1.95, 2.33, 3.55)
        cols.append(caixa(s, x + 0.1, 5.75, 2.25, 1.15, [(nome, 16, TEXTO, True), (desafio, 14, CINZA, False)],
                          espaco=1.05))
    ap.notas(s, "Estes são os cinco implantes do trabalho: a haste e o par articular da prótese da anca, o stent "
                "coronário, o scaffold para regenerar osso e o implante dentário. Cada um tem o problema crítico "
                "indicado pelo docente, e na haste acrescentei o stress shielding e a fadiga.")
    plano.append((s, f + cols, [], "conteudo"))

    # 3. Critérios-alvo
    s = novo()
    f = titulo(s, "Mais rígido não é melhor: o alvo é o osso", "A ideia", 3)
    v, d = video(s, "alvo_haste_diap", 0.6, 1.95, 8.6, "!!video_principal")  # 8,6 pol.: a largura do desenho
    lado = caixa(s, 9.45, 2.05, 3.3, 4.75, [
        ("Os métodos clássicos", 17, AZUL_TEXTO, True), ("só sabem que mais é melhor, ou que menos é melhor.", 16,
                                                         TEXTO, False),
        ("", 8, TEXTO, False),
        ("Os critérios-alvo", 17, AZUL_TEXTO, True), (f"dão a nota máxima a quem fica perto do valor ideal: na haste, "
                                                      f"o módulo do osso cortical ({ap.pt(n.extra['alvo'], 0)} GPa).",
                                                      16, TEXTO, False),
        ("", 8, TEXTO, False),
        ("Porquê?", 17, AZUL_TEXTO, True), ("Uma haste muito mais rígida do que o osso tira-lhe a carga, e o osso "
                                            "perde massa (stress shielding).", 16, TEXTO, False)], espaco=1.05)
    ap.notas(s, "Os métodos clássicos de decisão só sabem que mais é melhor ou que menos é melhor. Uma haste muito "
                "mais rígida do que o osso fica com a carga toda, o osso à volta deixa de ser estimulado e perde "
                "massa; chama-se stress shielding. Por isso usei critérios-alvo, "
                f"segundo Petković e colegas, 2025: a nota máxima vai para o material mais perto do valor ideal, aqui "
                f"o módulo do osso cortical, {ap.pt(n.extra['alvo'], 0)} GPa. O vídeo mostra cada liga no seu módulo e a "
                "nota que recebe.")
    plano.append((s, f + [lado], [(v, d, True)], "conteudo"))

    # 4. Funil
    s = novo()
    f = titulo(s, f"De {mil(fu['mp'])} materiais a um por dispositivo", "Os dados", 4)
    nv, nt = n_verificados()
    g = figura_caixa(s, grafico_funil(pd.read_csv(ap.RESULTS / "funil.csv"), CACHE / "funil.png"), 0.5, 1.95, 7.9,
                     4.9, nome="!!grafico")
    niveis = caixa(s, 8.75, 2.05, 4.0, 4.7, [
        ("1  Triagem em larga escala", 16, AZUL_TEXTO, True),
        (f"Materials Project (propriedades calculadas), {fu['elementos']} elementos com toxicidade de fontes "
         f"oficiais e {fu['normas']} normas ASTM de implantes.", 14, TEXTO, False),
        ("", 6, TEXTO, False),
        ("2  Filtros automáticos", 16, AZUL_TEXTO, True),
        ("Sem elementos tóxicos ou radioativos, estáveis, densidade e rigidez adequadas.", 14, TEXTO, False),
        ("", 6, TEXTO, False),
        ("3  Base verificada", 16, AZUL_TEXTO, True),
        (f"{fu['base']} materiais escolhidos à parte (não por filtro), com norma ou uso clínico; cada valor com "
         f"fonte e estado: {nv} dos {nt} verificados. É aqui que se decide.", 14, TEXTO, False)], espaco=1.05)
    ap.notas(s, f"A escolha começa larga: o Materials Project tem {mil(fu['mp'])} compostos inorgânicos com "
                "propriedades calculadas, que usei só para filtrar a toxicidade, a estabilidade e a rigidez: ficam "
                f"{mil(fu['nivel2'])}. A linha tracejada marca que a passagem para a base não é automática: os "
                f"{fu['base']} materiais da base foram escolhidos pelas normas ASTM e pelo uso clínico. Cada valor tem "
                f"fonte e estado: {nv} dos {nt} estão verificados, a começar pelos que mais pesam no resultado.")
    plano.append((s, f + [g, niveis], [], "conteudo"))

    # 5. O método em esquema: dados com fonte → critérios-alvo → três métodos → Monte Carlo → resposta
    s = novo()
    f = titulo(s, "Dos dados à resposta, em cinco passos", "O método", 5)
    passos = [("Dados com fonte", f"{nt} valores com mínimo, máximo, fonte e estado", CARTAO, AZUL_TEXTO),
              ("Critérios-alvo", "nota máxima perto do valor ideal", CARTAO, AZUL_TEXTO),
              ("Três métodos", "TOPSIS, WASPAS e VIKOR ordenam os materiais", CARTAO, AZUL_TEXTO),
              ("Monte Carlo", f"{mil(n.n_iter)} repetições com os dados a variar", QUENTE_FUNDO, QUENTE),
              ("Resposta", "o material e quanto confiar nele", NOITE, BRANCO)]
    wp, gap = 2.13, 0.37
    cs = []
    for i, (cab, txt, fundo_, cor) in enumerate(passos):
        x = 0.6 + i * (wp + gap)
        cartao(s, x, 1.95, wp, 1.75, fundo_)
        cs.append(caixa(s, x + 0.12, 2.05, wp - 0.24, 1.55, [(cab, 16, cor, True),
                                                           (txt, 14, SUB_NOITE if fundo_ == NOITE else TEXTO, False)],
                        ancora=MSO_ANCHOR.MIDDLE, espaco=1.05))
        if i < len(passos) - 1:
            caixa(s, x + wp - 0.02, 2.5, gap + 0.04, 0.6, "→", 24, CINZA, True, alinha=PP_ALIGN.CENTER)
    metodos = [("TOPSIS", "Escolhe o material mais perto do ideal e mais longe do pior."),
               ("WASPAS", "Soma e multiplica as notas: uma nota muito baixa pesa muito."),
               ("VIKOR", "Procura o compromisso e não deixa nenhum critério muito mau.")]
    for i, (m, frase) in enumerate(metodos):
        x = 0.6 + i * 4.13
        cartao(s, x, 4.0, 3.87, 1.35)
        cs.append(caixa(s, x + 0.25, 4.1, 3.4, 1.15, [(m, 18, AZUL_TEXTO, True), (frase, 14, TEXTO, False)],
                        espaco=1.05))
    mct = caixa(s, 0.6, 5.6, 12.13, 1.1, [
        ("Monte Carlo: os dados reais têm intervalos, não valores exatos. Repete-se a decisão "
         f"{mil(n.n_iter)} vezes, com valores sorteados dentro desses intervalos, e conta-se quantas vezes cada "
         "material fica em 1.º lugar.", 15, TEXTO, False)], espaco=1.05)
    ap.notas(s, "O método tem cinco passos. Primeiro, os dados: cada valor tem mínimo, máximo, fonte e estado. "
                "Segundo, os critérios-alvo, que dão a nota máxima ao valor ideal, como em Petković e colegas, 2025, cujos "
                "resultados reproduzi para verificar o código. Terceiro, três métodos ordenam os "
                "materiais: o TOPSIS escolhe o mais perto do ideal e mais "
                "longe do pior; o WASPAS soma e multiplica as notas, por isso castiga uma nota muito baixa; o VIKOR "
                "procura o compromisso. Quarto, como os dados têm intervalos, repito cada decisão "
                f"{mil(n.n_iter)} vezes com valores sorteados: é o Monte Carlo. No fim, a resposta: o material e "
                "quanto confiar nele.")
    plano.append((s, f + cs + [mct], [], "conteudo"))

    # 6 a 8. Casos quantitativos (mesmo esquema: Morph entre eles)
    def caso(nt, kicker, frase, render, video_nome, numeros_caso, nota):
        """video_nome: nome do vídeo em apresentacao/media/ ou, se for um Path, um gráfico estático no mesmo sítio."""
        s = novo()
        f = titulo(s, frase, kicker, nt)
        if isinstance(video_nome, Path):
            f.append(figura_caixa(s, video_nome, 0.6, 1.95, 8.6, 8.6 * 9 / 16, nome="!!video_principal"))
            vids = []
        else:
            v, d = video(s, video_nome, 0.6, 1.95, 8.6, "!!video_principal")
            vids = [(v, d, True)]
        im = figura_caixa(s, rend[render], 9.45, 1.85, 3.3, 1.7, nome="!!render_caso")
        nums = []
        for i, (val, leg) in enumerate(numeros_caso):
            y = 3.66 + i * 1.08
            cartao(s, 9.45, y, 3.3, 1.0)
            nums.append(caixa(s, 9.62, y + 0.04, 3.0, 0.92, [(val, 20, AZUL, True), (leg, 14, CINZA, False)],
                              ancora=MSO_ANCHOR.MIDDLE, nome=f"!!numero_{i}"))
        ap.notas(s, nota)
        plano.append((s, f + [im] + nums, vids, "caso"))

    hc = n.mc_haste
    caso(6, "Haste femoral", "Haste: a liga de titânio (Ti-6Al-4V) ganha quase sempre", "anca", "mc_haste_diap",
         [(f"{mil(hc['TOPSIS'])} a {mil(hc['VIKOR'])}", f"vitórias em {mil(n.n_iter)}, nos três métodos"),
          (f"C = {ap.pt(n.c_haste)}", "TOPSIS, cenário A"),
          (f"{round(n.extra['rigidez'])} ×", "mais rígido do que o alvo")],
         f"Na haste femoral, o Ti-6Al-4V ELI fica em 1.º com os três métodos e nos dois cenários, e no Monte Carlo "
         f"ganha em {mil(hc['TOPSIS'])} a {mil(hc['VIKOR'])} das {mil(n.n_iter)} repetições: é a escolha mais "
         f"sólida do trabalho; a fadiga das ligas vem dos ensaios de Okazaki, 2012, no mesmo laboratório. Mas o titânio continua cerca de {round(n.extra['rigidez'])} vezes mais rígido do que "
         f"o alvo. Pela relação de Gibson e Ashby, a mesma liga com {ap.pt(100 * n.extra['p_alvo'], 0)} % de "
         "porosidade chegaria ao alvo; como a resistência cai ainda mais, seria com porosidade graduada e o colo "
         "maciço. A solução para o stress shielding pode estar no desenho da haste, "
         "não noutra liga.")
    lo, hi = n.mc_stent_ptcr
    # Monte Carlo triangular do stent com o WASPAS: RELATORIO.md, secção 2.8 (results/mc_triangular.md)
    tri_l605, tri_ptcr = "52,7", "46,8"
    caso(7, "Stent coronário", FRASE_STENT_CURTA, "stent_arteria", "mc_stent_diap",
         [(f"ΔC = {ap.pt(n.dc_stent['A'])}", "cenário A; o VIKOR não os separa"),
          (f"ΔC = {ap.pt(n.dc_stent['B'])}", "cenário B: empate (ΔC < 0,01)"),
          (nb(f"{ap.pt(lo, 0)} a {ap.pt(hi, 0)} %"), "vitórias do Pt-Cr no Monte Carlo, consoante o método")],
         "No stent, no cenário A, o platina-crómio fica à "
         f"frente no TOPSIS, por {ap.pt(n.dc_stent['A'])}, e no VIKOR; o L605, no WASPAS. No cenário B "
         f"a margem é de {ap.pt(n.dc_stent['B'])}, abaixo da regra de empate. No Monte Carlo, o Pt-Cr ganha em "
         f"{ap.pt(lo, 0)} a {ap.pt(hi, 0)} % das repetições, mas três propriedades do Pt-Cr só estão "
         f"em fontes secundárias. Por isso: {FRASE_STENT}")
    sc = mc_primeiro("Scaffold")
    tcp, comp = sc[BTCP], sc[COMPOSITO]
    vsc = vencedores_ranking("Scaffold")
    n_comp = n_primeiro(vsc, COMPOSITO)
    res_comp, deg_comp = intervalo(COMPOSITO, "Resistência à compressão (scaffold)"), \
        intervalo(COMPOSITO, "Tempo de degradação")
    faixa = lambda d: nb(" | ".join(ap.pt(d[k], 0) for k in METODOS) + " %")  # noqa: E731
    caso(8, "Scaffold ósseo", FRASE_SCAFFOLD_CURTA, "scaffold",
         grafico_scaffold(vsc, sc, CACHE / "scaffold_carga.png"),
         [(nb(f"{n_comp} de {len(vsc)}"), "compósito em 1.º no ranking direto"),
          (faixa(tcp), "β-TCP em 1.º no Monte Carlo (TOPSIS | WASPAS | VIKOR)"),
          (f"ΔC = {ap.pt(n.dc_scaffold['A'])}", "β-TCP à frente só no TOPSIS, cenário A")],
         "No scaffold, no ranking direto, que conta também o suporte mecânico, o "
         f"compósito PCL com beta-TCP impresso em 3D fica em 1.º em {n_comp} das {len(vsc)} combinações; só no "
         f"TOPSIS do cenário A o beta-TCP poroso fica à frente, por {ap.pt(n.dc_scaffold['A'])}. Com os dados a "
         f"variar, o beta-TCP ganha mais vezes: em {ap.pt(tcp['TOPSIS'], 0)}, {ap.pt(tcp['WASPAS'], 0)} e "
         f"{ap.pt(tcp['VIKOR'], 0)} % das repetições. Verificar as fontes mudou duas vezes o "
         f"resultado: baixou a resistência do compósito para {ap.pt(res_comp[0], 1)} megapascal (Kawai, 2018) e "
         f"subiu a sua degradação para {ap.pt(deg_comp[0], 0)} meses (Zhong, 2026). O beta-TCP reabsorve-se em cerca "
         "de 13 meses (Hirata, 2006) e tem módulo de 546 megapascal (Montelongo, 2021). Por "
         f"isso: {FRASE_SCAFFOLD}")

    # 9. Espectro de confiança
    s = novo()
    f = titulo(s, "Nem todas as escolhas merecem a mesma confiança", "Os resultados", 9)
    niveis = confianca(n, True, simples=True)
    cor_nivel = {"Sólida": AZUL, "Indicativo": CINZA, "Muito próximos": QUENTE, "Empate técnico": QUENTE,
                 "Depende da leitura": QUENTE, "Depende da carga": QUENTE,
                 "Provável": AZUL_TEXTO}
    linhas = []
    for i, (disp, mat, nivel, frase) in enumerate(niveis):
        y = 2.0 + i * 0.95
        cartao(s, 0.6, y, 12.13, 0.82)
        chip = cartao(s, 0.8, y + 0.2, 2.1, 0.42, cor_nivel[nivel])
        tf = chip.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = nivel
        r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(14), True, LETRA, BRANCO
        linhas.append(caixa(s, 3.15, y + 0.04, 5.4, 0.76, [(disp, 16, TEXTO, True), (mat, 14, AZUL_TEXTO, True)],
                            ancora=MSO_ANCHOR.MIDDLE))
        linhas.append(caixa(s, 8.75, y + 0.06, 3.9, 0.72, [(frase, 14, CINZA, False)], ancora=MSO_ANCHOR.MIDDLE))
    ap.notas(s, "Junto os cinco resultados num espectro de confiança. A haste é sólida. "
                f"{FRASE_STENT_CURTA}. {FRASE_SCAFFOLD_CURTA}. No par articular e no implante dentário, a maior parte do peso cabe a escalas "
                "de 1 a 5, por isso o ranking é indicativo.")
    plano.append((s, f + linhas, [], "conteudo"))

    # 10. Resposta ao enunciado (Caso | Problema | Biomateriais a investigar | Proposta)
    s = novo()
    f = titulo(s, "Resposta ao enunciado: o que proponho em cada caso", "Síntese", 10)
    corpo = quadro(s, ("Caso", "Problema indicado", "Biomateriais comparados", "Proposta"), RESPOSTA,
                   LARG_RESPOSTA, h_cab=0.4, **ALT_RESPOSTA)
    ap.notas(s, "Este quadro responde ao enunciado: para cada caso, o problema indicado, os biomateriais que comparei "
                "e a proposta. Na haste, a liga de titânio numa haste não cimentada; no par articular, cerâmica contra "
                f"cerâmica; no stent, {sem_prefixo(FRASE_STENT_CURTA)}; no scaffold, depende da carga; no implante "
                "dentário, titânio de grau 4.")
    plano.append((s, f + corpo, [], "conteudo"))

    # 11. Como caracterizar e testar (questões 9 a 11 do enunciado)
    s = novo()
    f = titulo(s, "Como caracterizar e testar cada proposta", "Fabrico, testes e caracterização (questões 9 a 11)",
               11)
    corpo = quadro(s, ("Proposta", "Fabrico", "Testes antes do uso clínico", "Caracterização"), FABRICO,
                   (2.2, 3.45, 3.3, 3.18), y0=1.6, h_cab=0.4, h_linha=0.86)
    ap.notas(s, "Antes do uso clínico, cada proposta tem de ser fabricada, testada e caracterizada. Todas passam pela "
                "série ISO 10993 de biocompatibilidade, e cada uma tem a sua norma mecânica, como a ISO 7206-4 na "
                "fadiga da haste e a ASTM F2477 na do stent. Na caracterização entram a tração, a "
                "microscopia eletrónica e, no scaffold, a micro-CT da porosidade.")
    plano.append((s, f + corpo, [], "conteudo"))

    # 12. Perfis de doente
    s = novo()
    f = titulo(s, "E o doente? Uma condição clínica pode mudar a escolha", "Extensão: perfis de doente", 12)
    exemplos = perfis_exemplo()
    cs = []
    exemplos = exemplos[:3]
    wc = (12.13 - 0.3 * (len(exemplos) - 1)) / max(1, len(exemplos))
    for i, (desc, caso_, antes, depois) in enumerate(exemplos):
        x = 0.6 + i * (wc + 0.3)
        cartao(s, x, 2.05, wc, 3.15)
        cs.append(caixa(s, x + 0.35, 2.25, wc - 0.7, 2.9, [
            (caso_, 14, AZUL_TEXTO, True), (desc, 22, TEXTO, True), ("", 10, TEXTO, False),
            ("sem perfil", 14, CINZA, False), (ns.com_codigo(antes), 18, CINZA, False), ("", 6, TEXTO, False),
            ("com o perfil", 14, CINZA, False), (ns.com_codigo(depois), 22, QUENTE, True)], espaco=1.02))
    desc_h1, _, h1_topsis, _ = sensibilidade_perfil("H1")
    regras = caixa(s, 0.6, 5.45, 8.6, 1.4, [
        (f"{n_regras()} regras com fonte por regra", 17, TEXTO, True),
        ("Excluem materiais (alergia confirmada) ou reforçam critérios (osteoporose, diabetes, tabagismo, "
         "radioterapia, RM frequente, entre outras).", 14, CINZA, False),
        (f"{desc_h1}: o TOPSIS passa a {ns.com_codigo(h1_topsis)}; é sensibilidade ao fator do perfil, não "
         "recomendação.", 14, CINZA, False)], espaco=1.05)
    aviso = cartao(s, 9.45, 5.65, 3.28, 0.95, QUENTE_FUNDO)
    av = caixa(s, 9.6, 5.72, 3.0, 0.85, [(AVISO, 14, TEXTO, True)], ancora=MSO_ANCHOR.MIDDLE)
    ap.notas(s, "A última camada é o doente: uma tabela de regras, cada uma com a sua fonte, exclui materiais numa "
                "alergia confirmada ou reforça critérios em condições como a osteoporose ou a diabetes. Nos exemplos, "
                "um defeito ósseo sem carga troca o compósito pelo beta-TCP, e uma alergia ao titânio confirmada por "
                "teste troca o titânio pela zircónia. No jovem ativo, o TOPSIS passa ao cobalto-crómio, "
                "como teste de sensibilidade. Não é ferramenta clínica nem recomendação "
                "terapêutica.")
    plano.append((s, f + cs + [regras, aviso, av], [], "conteudo"))

    # 13. Impacto no doente: haste (FEA) e stent (CFD) num só diapositivo; os mapas ficam na reserva (18 e 19)
    s = novo()
    f = titulo(s, "Impacto no doente: haste e stent", "Extensão exploratória: simulações; hipóteses no relatório",
               13)
    corpo = []
    rz = razoes_fea() if FEA_CSV.exists() else {}
    dcfd = dados_cfd() if CFD_JSON.exists() else {}
    for x, w, cab in ((0.6, 6.2, "Haste: elementos finitos no fémur"), (7.1, 5.65, "Stent: sangue numa coronária")):
        corpo.append(caixa(s, x, 1.62, w, 0.4, cab, 18, AZUL_TEXTO, True))
    if rz:
        corpo.append(figura_caixa(s, grafico_fea(rz, CACHE / "fea_razoes_juntas.png", 6.2, 3.9), 0.5, 2.1, 6.2,
                                  3.9, nome="!!grafico_fea"))
    if dcfd:
        corpo.append(figura_caixa(s, grafico_cfd(dcfd, CACHE / "cfd_area_baixa_juntas.png", 5.65, 3.9), 7.1, 2.1,
                                  5.65, 3.9, nome="!!grafico_cfd"))
    corpo.append(caixa(s, 0.6, 6.1, 6.2, 0.75, "Quanto mais rígida a haste, menos carga chega ao osso junto ao "
                       "colo.", 16, TEXTO, True))
    corpo.append(caixa(s, 7.1, 6.1, 5.65, 0.75, "Struts mais grossos deixam mais parede com o sangue a passar "
                       "devagar.", 16, TEXTO, True))
    nota = "Duas simulações exploratórias. Na haste, simulei o fémur com a carga do pico da marcha"
    if all(m in rz for m in FEA_ORDEM):
        nota += (f": junto ao colo, a energia no osso, em relação ao osso sem prótese, fica em "
                 f"{ap.pt(rz['Co-Cr-Mo'][1], 2)} com o cobalto-crómio, {ap.pt(rz['316L'][1], 2)} com o aço e "
                 f"{ap.pt(rz['Ti-6Al-4V'][1], 2)} com a liga de titânio; a liga porosa chega a "
                 f"{ap.pt(rz['Ti-6Al-4V poroso'][1], 1)}")
    nota += ". No stent, simulei o sangue a pulsar numa coronária de 3 milímetros"
    p_, a_ = dcfd.get("Pt-Cr"), dcfd.get("316L")
    if p_ and a_:
        nota += (f": com struts de {p_['espessura_um']:.0f} micrómetros, {100 * p_['fracao_baixa']:.0f} % da "
                 f"parede entre os struts tem fluxo lento, abaixo de 0,4 pascal, um limiar operacional deste "
                 f"estudo; com {a_['espessura_um']:.0f}, "
                 f"{100 * a_['fracao_baixa']:.0f} %. O efeito é da espessura, não da liga")
    ap.notas(s, nota + ". O micromovimento máximo na interface é de 66 micrómetros, abaixo dos 150 a partir dos "
                "quais a fixação deixa de ser por osso, em cães (Pilliar, 1986). A energia é um indicador mecânico de redistribuição de carga, não previsão de perda "
                "óssea.")
    plano.append((s, f + corpo, [], "conteudo"))

    # 14. Duas ferramentas, a mesma resposta (SOLIDWORKS e Ansys Mechanical; OpenFOAM e Ansys Fluent)
    s = novo()
    f = titulo(s, "Duas ferramentas, a mesma resposta", "Verificação numérica cruzada das simulações", 14)
    corpo = []
    if SW_JSON.exists() and MAPDL_JSON.exists() and FLUENT_JSON.exists():
        r2, a2 = razoes_duas_ferramentas(), areas_duas_ferramentas()
        dmax = dif_max_razoes(r2)
        corpo.append(figura_caixa(s, grafico_duas_ferramentas(r2, a2, CACHE / "duas_ferramentas.png"), 0.6, 1.7,
                                  12.13, 3.6, nome="!!grafico_duas"))
        ordem = " > ".join(NOMES_FEA_DUAS[m] for m in sorted(r2, key=lambda m: -r2[m]["Ansys"][1]))
        pc = lambda li, k: nb(f"{ap.pt(a2[li][k], 1)} %")  # noqa: E731
        corpo.append(caixa(s, 0.6, 5.42, 6.6, 1.3, [
            ("Haste: SOLIDWORKS e Ansys Mechanical", 16, AZUL_TEXTO, True),
            (f"A mesma proteção de tensões no osso (razões em Gruen 1 e 7 com {ap.pt(dmax, 1)} % de diferença ou "
             f"menos) e a mesma ordem: {ordem}.", 14, TEXTO, False)], espaco=1.02))
        corpo.append(caixa(s, 7.45, 5.42, 5.3, 1.3, [
            ("Stent: OpenFOAM e Ansys Fluent", 16, AZUL_TEXTO, True),
            (f"A mesma área com tensão de corte baixa (abaixo de {nb('0,4 Pa')}): Pt-Cr {pc('Pt-Cr', 'OpenFOAM')} contra "
             f"{pc('Pt-Cr', 'Ansys Fluent')}; 316L {pc('316L', 'OpenFOAM')} contra {pc('316L', 'Ansys Fluent')}.",
             14, TEXTO, False)], espaco=1.02))
        nota2 = ("Repeti as simulações num segundo programa, independente. "
                 f"Na haste, o SOLIDWORKS e o Ansys Mechanical dão a mesma proteção de tensões, com diferenças de "
                 f"{ap.pt(dmax, 1)} % ou menos, e a mesma ordem dos materiais. No stent, o "
                 f"OpenFOAM e o Ansys Fluent dão a mesma área com fluxo lento: {ap.pt(a2['Pt-Cr']['OpenFOAM'], 1)} "
                 f"contra {ap.pt(a2['Pt-Cr']['Ansys Fluent'], 1)} % com os struts finos. É uma verificação "
                 "numérica cruzada, não uma validação experimental.")
    else:
        nota2 = "Comparação entre ferramentas por gerar (results/fea_haste_mapdl.json e results/cfd_fluent.json)."
    corpo.append(caixa(s, 0.6, 6.68, 11.4, 0.32, "Verificação numérica cruzada, não validação experimental: mesma "
                       "geometria, condições e (no stent) malha.", 14, CINZA))
    ap.notas(s, nota2)
    plano.append((s, f + corpo, [], "conteudo"))

    # 15. Conclusão (com o endereço da versão web e o QR)
    s = novo(True)
    q = caixa(s, 0.6, 0.9, 8.0, 2.2, "Um ranking só vale o que valem os dados que o alimentam", 38, BRANCO, True,
              nome="!!titulo")
    pts = caixa(s, 0.6, 3.2, 7.6, 2.7, [
        ("O ideal depende do corpo: os critérios-alvo dão nota máxima ao meio-termo certo.", 18, BRANCO, False),
        ("", 8, BRANCO, False),
        ("Medir a confiança separa escolhas sólidas (haste) de ligas muito próximas (stent).", 18, BRANCO, False),
        ("", 8, BRANCO, False),
        ("Verificar as fontes mudou o scaffold: a escolha passou a depender da carga.", 18, BRANCO,
         False)], espaco=1.1)
    fim = caixa(s, 0.6, 5.95, 7.6, 1.2, [("Obrigado.", 22, AZUL, True),
                                         (f"Ao {ORIENTADOR}, docente de Próteses e Órgãos Artificiais, pela "
                                          "orientação.", 14, SUB_NOITE, False)])
    figura_caixa(s, rend["dentario"], 9.0, 0.5, 3.75, 3.9)
    entradas_fim = [q, pts, fim]
    if QR_WEB.exists():
        qr_fundo = cartao(s, 10.85, 4.6, 1.9, 1.9, BRANCO)
        figura(s, QR_WEB, 10.93, 4.68, 1.74, 1.74, nome="!!qr_web")
        entradas_fim += [qr_fundo]
    dominio, caminho = URL_WEB.split("/", 1)
    web = caixa(s, 8.0, 4.6, 2.7, 1.9, [("Experimente no navegador", 14, BRANCO, True),
                                        (dominio + "/", 14, SUB_NOITE, False), (caminho, 14, SUB_NOITE, False)],
                alinha=PP_ALIGN.RIGHT, ancora=MSO_ANCHOR.MIDDLE, nome="!!url_web")
    entradas_fim.append(web)
    if logo:
        entradas_fim.append(figura(s, logo, 0.6, 0.2, h=0.6, nome="!!logo_projeto"))
    ap.notas(s, "Fecho com três ideias. Primeira: o ideal depende do corpo; os critérios-alvo dão a nota máxima ao "
                "meio-termo certo, como a rigidez do osso. Segunda: medir a confiança separa escolhas sólidas, como a "
                "haste, de materiais muito próximos, como no stent. Terceira, a mais importante: um ranking só vale o "
                "que valem os dados; verificar as fontes mudou o scaffold, que passou a depender da carga. O próximo passo é ligar a "
                "ferramenta a dados reais de doentes. Obrigado.")
    plano.append((s, entradas_fim, [], "capa"))

    # Reserva
    s = novo(True)
    t = caixa(s, 0.6, 2.8, 12.1, 1.2, "Diapositivos de reserva", 40, BRANCO, True, nome="!!titulo")
    sub = caixa(s, 0.6, 3.9, 12.1, 0.8, "A aplicação, as simulações e porquê cada ferramenta, em linguagem simples",
                20, SUB_NOITE)
    ap.notas(s, "Diapositivos de reserva, para as perguntas.")
    plano.append((s, [t, sub], [], "capa"))

    def reserva(nt, frase, kicker, blocos, nota, imagem=None):
        s = novo()
        f = titulo(s, frase, kicker, nt)
        extra = []
        if imagem:
            corpo = caixa(s, 0.6, 2.05, 6.3, 4.7, blocos, espaco=1.05)
            extra.append(figura_caixa(s, imagem, 7.2, 2.0, 5.55, 4.7))
        else:  # pares (cabeçalho, texto) em cartões, duas colunas
            pares = [blocos[i:i + 2] for i in range(0, len(blocos), 2)]
            linhas_ = (len(pares) + 1) // 2
            hc = min(2.6, (4.7 - 0.3 * (linhas_ - 1)) / linhas_)
            caixas_ = []
            for i, (cab, corpo_) in enumerate(pares):
                x, y = 0.6 + (i % 2) * 6.21, 2.05 + (i // 2) * (hc + 0.3)
                cartao(s, x, y, 5.92, hc)
                caixas_.append(caixa(s, x + 0.3, y + 0.25, 5.32, hc - 0.4,
                                     [(cab[0], 20, AZUL_TEXTO, True), (corpo_[0], 16, TEXTO, False)], espaco=1.08))
            corpo = caixas_[0]
            extra += caixas_[1:]
        ap.notas(s, nota)
        plano.append((s, f + [corpo] + extra, [], "conteudo"))

    # 17. Aplicação (reserva): vídeo da demonstração quando existir; senão, um marcador claro
    s = novo()
    f = titulo(s, "Uma aplicação para explorar as escolhas", "Reserva: a aplicação", 17)
    demo = primeiro_existente(DEMO_APP) if demo else None
    vids = []
    if demo:
        import imageio_ffmpeg
        _, seg = imageio_ffmpeg.count_frames_and_secs(str(demo))
        capa = capa_video(demo, CACHE / f"{demo.stem}_capa.png")
        v = s.shapes.add_movie(str(demo), Inches(0.6), Inches(1.95), Inches(8.4), Inches(4.725),
                               poster_frame_image=str(capa), mime_type="video/mp4")
        v.name = "!!video_principal"
        vids.append((v, int(seg * 1000), False))  # reserva: o vídeo corre ao clicar
        corpo = [v]
    else:
        m = cartao(s, 0.6, 1.95, 8.4, 4.725, NOITE, nome="!!video_principal")
        corpo = [m, caixa(s, 1.0, 3.3, 7.6, 2.0, [
            ("[VÍDEO DA DEMONSTRAÇÃO POR INSERIR]", 20, BRANCO, True),
            ("Gerar app-nativa/demo/demo_app_60s.mp4 e exportar de novo a apresentação.", 14, SUB_NOITE, False)],
            alinha=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE, nome="!!marcador_demo")]
    lista = caixa(s, 9.35, 1.95, 3.4, 3.3, [
        ("Ranking e Monte Carlo", 14, AZUL_TEXTO, True), ("pesos ao vivo; 10 000 repetições no núcleo Rust", 12, TEXTO, False),
        ("Doente em tempo real", 14, AZUL_TEXTO, True), ("regras, pesos e exclusões, com a justificação", 12, TEXTO, False),
        ("Comparar materiais", 14, AZUL_TEXTO, True), ("dois a dois, com a fonte e o estado", 12, TEXTO, False),
        ("Simulação", 14, AZUL_TEXTO, True), ("FEA da haste (von Mises) e CFD do stent", 12, TEXTO, False),
        ("Histórico e PDF", 14, AZUL_TEXTO, True), ("guardar e exportar cada análise", 12, TEXTO, False),
        ("Modo apresentação", 14, AZUL_TEXTO, True), ("vista limpa, letra e gráficos grandes", 12, TEXTO, False)],
        espaco=1.0)
    dom, cam = URL_WEB.split("/", 1)
    web = caixa(s, 9.35, 5.3, 3.4, 1.3, [("Aplicação para Windows e no navegador:", 14, CINZA, False),
                                         (dom + "/", 14, AZUL_TEXTO, True), (cam, 14, AZUL_TEXTO, True)],
                espaco=1.02)
    ap.notas(s, "Para explorar estas escolhas fiz uma aplicação, para Windows e no navegador. Escolhe-se o caso, "
                "ajustam-se os pesos e o ranking dos três métodos reage na hora; um botão corre o Monte Carlo no "
                "núcleo, escrito em Rust e testado contra o código Python. No ecrã do doente, um painel mostra em "
                "tempo real que regras se aplicam, como mudam os pesos e que materiais saem, sempre com a "
                "justificação. Pode comparar-se dois materiais, ver as simulações (mapas de von Mises da haste e a "
                "CFD do stent), guardar cada análise no histórico, exportá-la em PDF e passar ao modo apresentação. "
                "É uma ferramenta educativa, não clínica.")
    plano.append((s, f + corpo + [lista, web], vids, "conteudo"))

    # 18. Impacto no doente: a haste (FEA), com o mapa (reserva)
    s = novo()
    f = titulo(s, "Impacto no doente: a haste", "Reserva: simulação da haste", 18)
    corpo, vids = video_simulacao(s, "sim_fea_haste", LEGENDA_FEA)
    if rz:
        if not corpo:
            corpo.append(figura_caixa(s, FEA_MAPA if FEA_MAPA.exists() else FEA_MAPA.with_name("sed_corte_coronal.png"),
                                      0.5, 1.75, 6.9, 4.3, nome="!!mapa_fea"))
        corpo.append(figura_caixa(s, grafico_fea(rz, CACHE / "fea_razoes.png"), 7.55, 1.75, 5.25, 4.3,
                                  nome="!!grafico_fea"))
    frase = caixa(s, 0.6, 6.15, 9.4, 0.5, "Quanto mais rígida a haste, menos carga chega ao osso junto ao colo.", 18,
                  TEXTO, True)
    rot = caixa(s, 10.0, 6.15, 2.75, 0.62, "Simulação exploratória; hipóteses no relatório", 14, CINZA,
                alinha=PP_ALIGN.RIGHT)
    ap.notas(s, "O mapa mostra a energia de deformação no osso, num corte do fémur, sem prótese e com duas hastes. "
                "No SOLIDWORKS e no Ansys a ordem é a mesma; junto ao colo, em Gruen 7, os valores das ligas maciças "
                "são cerca do dobro, porque o fémur é simplificado: só a ordem é robusta. É um indicador mecânico de "
                "redistribuição de carga, não uma previsão de perda óssea nem de descolamento; a interface colada não "
                "modela micromovimento. A haste porosa exigiria porosidade graduada, com o colo maciço.")
    plano.append((s, f + corpo + [frase, rot], vids, "conteudo"))

    # 19. Impacto no doente: a espessura do strut (CFD), com o mapa (reserva)
    s = novo()
    f = titulo(s, "Impacto no doente: a espessura do strut", "Reserva: simulação do stent", 19)
    corpo, vids = video_simulacao(s, "sim_cfd_stent", LEGENDA_CFD)
    if dcfd and CFD_MAPA.exists():
        if not corpo:
            corpo.append(figura_caixa(s, CFD_MAPA, 0.5, 1.75, 6.9, 4.3, nome="!!mapa_cfd"))
        corpo.append(figura_caixa(s, grafico_cfd(dcfd, CACHE / "cfd_area_baixa.png"), 7.55, 1.75, 5.25, 4.3,
                                  nome="!!grafico_cfd"))
    frase = caixa(s, 0.6, 6.15, 9.4, 0.5, "Struts mais grossos deixam mais parede com o sangue a passar devagar.", 18,
                  TEXTO, True)
    rot = caixa(s, 10.0, 6.15, 2.75, 0.62, "Simulação exploratória; hipóteses no relatório", 14, CINZA,
                alinha=PP_ALIGN.RIGHT)
    ap.notas(s, "O mapa é a parede da artéria aberta como uma folha: a laranja, onde o sangue roça devagar, abaixo de "
                "0,4 pascal, um limiar operacional deste estudo; a geometria é idealizada, sem lesão e com parede "
                "rígida. O efeito é da espessura, não da liga: a espessura é do dispositivo. É um indicador do "
                "fluxo, não uma previsão de reestenose. A fração depende do limiar: a partir de cerca de 0,8 pascal "
                "as duas espessuras ficam praticamente iguais; a tensão média relativa à do tubo sem stent, 0,47 "
                "contra 0,27, não depende dele.")
    plano.append((s, f + corpo + [frase, rot], vids, "conteudo"))

    lo316, hi316, u316 = intervalo("Aço inox 316L", "Módulo de Young")
    reserva(20, "Porquê Monte Carlo? Porque os dados têm intervalos", "Reserva: estatística", [
        ("A ideia", 17, AZUL_TEXTO, True),
        (f"Um valor como o módulo de uma liga vem num intervalo (por exemplo {ap.pt(lo316, 0)} a {ap.pt(hi316, 0)} "
         f"{u316} no aço 316L). Em vez de "
         "usar só o meio, sorteia-se um valor dentro do intervalo, decide-se, e repete-se.", 15, TEXTO, False),
        ("Exemplo", 17, AZUL_TEXTO, True),
        (f"Na haste, em {mil(n.n_iter)} repetições o Ti-6Al-4V ficou em 1.º em {mil(hc['TOPSIS'])} "
         "(TOPSIS). Isto não é uma probabilidade clínica: é a estabilidade da decisão face aos dados.", 15, TEXTO,
         False),
        ("Também testado", 17, AZUL_TEXTO, True),
        ("Distribuição triangular em vez de uniforme, pesos a variar, pesos objetivos e cenários.", 15, TEXTO,
         False)],
            "O Monte Carlo serve para uma pergunta simples: a escolha mantém-se quando os dados variam dentro do que "
            "as fontes dizem? Sorteia-se cada valor dentro do seu intervalo, decide-se de novo e conta-se. A "
            "frequência mede a estabilidade da decisão, não a probabilidade de sucesso clínico. O gráfico é o do "
            "stent: o Pt-Cr ganha na maioria das repetições com os três métodos, mas não em todas. Com distribuições "
            f"triangulares, os vencedores mantêm-se, exceto no stent com o WASPAS: o L605 passa a {tri_l605} % "
            f"contra {tri_ptcr} % do Pt-Cr.",
            grafico_mc_stent(mc_primeiro("Stent"), CACHE / "mc_stent.png"))
    reserva(21, "Porquê três métodos? Porque discordam de forma útil", "Reserva: os algoritmos", [
        ("TOPSIS", 17, AZUL_TEXTO, True),
        ("Calcula a distância de cada material ao ideal (o melhor em tudo) e ao pior; C perto de 1 é bom. Um ponto "
         "forte pode compensar um fraco.", 15, TEXTO, False),
        ("WASPAS", 17, AZUL_TEXTO, True),
        ("Média de uma soma e de um produto das notas: no produto, uma nota quase nula puxa tudo para baixo.", 15,
         TEXTO, False),
        ("VIKOR", 17, AZUL_TEXTO, True),
        ("Pesa o desempenho médio e o pior critério de cada material e procura o compromisso.", 15, TEXTO, False),
        ("Borda", 17, AZUL_TEXTO, True),
        ("Cada método dá pontos pela posição (1.º de 5 vale 4 pontos); somam-se os pontos dos três.", 15, TEXTO,
         False)],
            "Usei três métodos porque discordam de forma útil. O TOPSIS deixa compensar; o WASPAS castiga notas "
            "muito baixas; o VIKOR olha para o pior critério. Quando os três concordam, como na haste, a escolha é "
            "forte; quando discordam, como no stent, é sinal de que os materiais estão muito próximos. O Borda junta as três ordens.")
    # 22. Outros métodos (extensão, ramo extensoes: methods/extra.py, sobol_pesos.py): só o que fazem, sem resultados
    reserva(22, "Outros métodos para pôr a escolha à prova", "Reserva: métodos adicionais", [
        ("Borda", 17, AZUL_TEXTO, True),
        ("Junta as ordens dos três métodos numa só: cada posição vale pontos e somam-se.", 15, TEXTO, False),
        ("Sobol", 17, AZUL_TEXTO, True),
        ("Faz variar todos os pesos ao mesmo tempo e mede quais mudam mais o resultado.", 15, TEXTO, False),
        ("PROMETHEE II", 17, AZUL_TEXTO, True),
        ("Compara os materiais dois a dois, critério a critério, e soma as vantagens de cada um.", 15, TEXTO,
         False),
        ("MOORA", 17, AZUL_TEXTO, True),
        ("Põe as notas na mesma escala, soma as boas e subtrai as más.", 15, TEXTO, False),
        ("CODAS", 17, AZUL_TEXTO, True),
        ("Mede quão longe cada material fica do pior, com uma segunda distância para desempatar.", 15, TEXTO,
         False)],
            "Além dos três métodos principais, há outros para pôr a escolha à prova. O Borda junta as três ordens "
            "numa só. A análise de Sobol faz variar todos os pesos ao mesmo tempo e diz quais pesam mais no "
            "resultado. O PROMETHEE II, o MOORA e o CODAS são três métodos de decisão diferentes, para ver se a "
            "ordem depende do método. O Borda está no trabalho principal; os outros quatro são uma extensão, fora "
            "do relatório.")
    reserva(23, "A rigidez certa pode vir do desenho, não só do material", "Reserva: Gibson e Ashby", [
        ("A relação", 17, AZUL_TEXTO, True),
        ("Num material poroso, o módulo cai com a porosidade: " + nb("E = Es (1 − p)²") + ", em que Es é o módulo "
         "da liga maciça.", 15, TEXTO, False),
        ("Na haste", 17, AZUL_TEXTO, True),
        (f"O Ti-6Al-4V maciço é cerca de {round(n.extra['rigidez'])} vezes mais rígido do que o alvo; com cerca de "
         f"{ap.pt(100 * n.extra['p_alvo'], 0)} % de porosidade chegaria aos {ap.pt(n.extra['alvo'], 0)} GPa.", 15,
         TEXTO, False),
        ("Limite", 17, AZUL_TEXTO, True),
        ("A resistência fica numa fração ainda menor do que o módulo, e a fadiga do titânio poroso é sensível aos "
         "defeitos: na prática, porosidade graduada, com o colo e a zona de maior tensão maciços.", 15, TEXTO, False)],
            "A relação de Gibson e Ashby diz que o módulo de um material poroso cai com o quadrado da fração sólida. "
            "Serve para mostrar que o stress shielding se pode atacar pelo desenho, com titânio poroso, e não só "
            "trocando de liga. Mas a resistência e a fadiga ficam numa fração ainda menor do que o módulo: a 60 % de "
            "porosidade, a resistência fica em cerca de 8 % da liga maciça e o módulo em 16 %. Por isso os desenhos "
            "usam porosidade graduada, com o colo maciço.",
            grafico_gibson(CACHE / "gibson_ashby.png"))
    reserva(24, "O modelo reproduz o artigo de referência", "Reserva: verificação da implementação", [
        ("Reprodução", 17, AZUL_TEXTO, True),
        ("O caso da prótese da anca de Petković et al. (2025) foi reproduzido: os 15 coeficientes C publicados, com "
         "diferença máxima de cerca de 5 × 10⁻⁶; o WASPAS e o VIKOR dos casos 1 e 2 também batem (abaixo de 10⁻⁵).",
         15, TEXTO, False),
        ("Pormenor", 17, AZUL_TEXTO, True),
        ("Uma célula da tabela impressa do artigo difere do valor que reproduz os resultados publicados; ficou "
         "documentada no relatório.", 15, TEXTO, False)],
            "Antes de usar o código nos meus casos, reproduzi o caso publicado por Petković e colegas: os quinze "
            "coeficientes batem certo. Os pormenores da diferença numa célula estão no relatório.")
    reserva(25, "Os dados vêm de fontes abertas e verificadas", "Reserva: fontes e licenças", [
        ("Materials Project", 17, AZUL_TEXTO, True),
        ("Dados abertos (licença CC BY 4.0), propriedades calculadas por DFT; só para filtrar. Jain et al., APL "
         "Materials, 2013.", 15, TEXTO, False),
        ("Toxicidade dos elementos", 17, AZUL_TEXTO, True),
        ("ICH Q3D(R2), classes das impurezas elementares; monografias da IARC; elementos sem isótopos estáveis.", 15,
         TEXTO, False),
        ("Normas e base verificada", 17, AZUL_TEXTO, True),
        (f"{fu['normas']} normas ASTM de materiais para implantes; cada valor da base com mínimo, máximo, típico, "
         "fonte e estado.", 15, TEXTO, False)],
            "As fontes da triagem são abertas: o Materials Project sob licença CC BY 4.0, a guideline ICH Q3D e as "
            "monografias da IARC para a toxicidade, e as normas ASTM. Os valores que entram na decisão vêm só da base "
            "verificada.")

    massa_csv = ap.RESULTS / "massa_haste.csv"
    if massa_csv.exists():
        mh = pd.read_csv(massa_csv).set_index("material")["massa_g"]
        dm = mh.max() - mh.min()
        reserva(26, "A massa da haste não pesa; os iões contam", "Reserva: impacto no doente", [
            ("Massa da haste (a mesma forma)", 17, AZUL_TEXTO, True),
            (f"Cobalto-crómio {mh['Co-Cr-Mo forjado']:.0f} g, aço {mh['Aço inox 316L']:.0f} g, liga de titânio "
             f"{mh['Ti-6Al-4V ELI']:.0f} g: {dm:.0f} g de diferença.", 15, TEXTO, False),
            ("Centro de massa do corpo", 17, AZUL_TEXTO, True),
            # results/massa_haste.md e RELATORIO.md, secção 8.5 (Orendurff et al., 2004)
            ("Muda no máximo 0,6 mm: desprezável. A cada passo, o centro de massa sobe e desce 2,7 a 4,8 cm.", 15,
             TEXTO, False),
            ("Iões de cobalto e crómio", 17, AZUL_TEXTO, True),
            ("Raramente, com níveis muito altos no sangue: coração, tiroide, audição e visão. No Reino Unido, a "
             "partir de 7 µg/L há seguimento apertado.", 15, TEXTO, False),
            ("Níquel", 17, AZUL_TEXTO, True),
            ("Todas as ligas permanentes do stent têm níquel (9 a 35 %); a alergia ao níquel afeta 8 a 19 % dos "
             "adultos europeus.", 15, TEXTO, False)],
                f"Duas perguntas que costumam aparecer. A massa: a haste de cobalto-crómio pesa cerca de "
                f"{mh['Co-Cr-Mo forjado']:.0f} gramas e a de titânio cerca de {mh['Ti-6Al-4V ELI']:.0f}; trocar uma "
                "pela outra muda o centro de massa do corpo no máximo 0,6 milímetros, o que é desprezável. Os iões: o "
                "cobalto e o crómio só dão efeitos no corpo todo raramente, com níveis muito altos, sobretudo com "
                "desgaste de pares metal-metal; o níquel conta no stent, porque todas as ligas permanentes o têm, mas a "
                "relevância clínica da alergia ao níquel em stents é debatida (Köster 2000 contra Norgaz 2005), por "
                "isso o perfil de doente com alergia ao níquel é exploratório. Fontes: Bradberry 2014, MHRA 2017, "
                "Ahlström 2019, Orendurff 2004.")

    for k, (s, entradas, vids, esquema) in enumerate(plano):
        morph = k > 0 and plano[k - 1][3] == esquema == "caso"
        if esquema == "abertura":  # o vídeo começa logo e o resto entra quando ele acaba
            ap.aplicar_animacao(s, entradas, vids, morph, atraso_entradas=vids[0][1], atraso_video=0)
        else:
            ap.aplicar_animacao(s, entradas, vids, morph)
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(destino)
    return destino


def exportar(pptx: Path = DESTINO) -> None:
    """No PowerPoint: guarda com as letras embutidas, exporta o PDF, o plano B e as imagens (1920 × 1080)."""
    import shutil

    import win32com.client
    app = win32com.client.DispatchEx("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(pptx.resolve()), False, False, False)
        pres.SaveAs(str(pptx.resolve()), 24, True)  # 24 = ppSaveAsOpenXMLPresentation; True = embutir letras
        pres.SaveAs(str(pptx.with_suffix(".pdf").resolve()), 32)
        pasta = APRES / "imagens_b"
        pasta.mkdir(exist_ok=True)
        for i in range(1, pres.Slides.Count + 1):
            pres.Slides(i).Export(str(pasta / f"diapositivo_{i:02d}.png"), "PNG", 1920, 1080)
        pres.Close()
    finally:
        app.Quit()
    shutil.copyfile(pptx.with_suffix(".pdf"), APRES / "apresentacao_plano_B.pdf")


def main() -> None:
    a = argparse.ArgumentParser()
    a.add_argument("--exportar", action="store_true")
    a.add_argument("--renders", type=Path, default=None)
    a.add_argument("--sem-demo", action="store_true", help="deixa o marcador no lugar do vídeo da demo")
    args = a.parse_args()
    p = construir(renders=args.renders, demo=not args.sem_demo)
    print(p.relative_to(RAIZ))
    import imageio_ffmpeg
    pal = palavras_notas(Presentation(str(p)))
    _, seg = imageio_ffmpeg.count_frames_and_secs(str(ap.MEDIA / f"{ABERTURA}.mp4"))
    print(f"notas dos {N_PRINCIPAIS} principais: {pal} palavras; cerca de {minutos_fala(pal, seg):.1f} min "
          f"a {PALAVRAS_POR_MINUTO} palavras por minuto, com o vídeo de abertura ({seg:.0f} s)")
    if not primeiro_existente(LOGO_PROJETO):
        print("logótipo do projeto: ainda não existe (" + ", ".join(map(str, LOGO_PROJETO)) + ")")
    if args.sem_demo or not primeiro_existente(DEMO_APP):
        print("vídeo da demonstração: não incluído; o diapositivo 17 leva um marcador")
    if args.exportar:
        exportar(p)


if __name__ == "__main__":
    main()
