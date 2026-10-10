"""Gera a versão Word do relatório (relatorio/RELATORIO_vN.docx) com o mesmo formato do PDF.

Passos:
1. Cria o modelo relatorio/formato/referencia.docx (a partir do modelo do pandoc): Times New Roman 11 pt,
   títulos a negrito (14 pt, 12 pt e 11 pt), margens de 2 cm, A4, texto justificado, espaçamento 1,15,
   número de página no rodapé e os estilos da primeira página, das legendas e das referências.
2. Converte relatorio/RELATORIO_word.md (gerado por scripts/montar_relatorio.py) com o pandoc
   (--reference-doc); as equações em LaTeX passam a equações do Word.
3. Acerta o docx com python-docx: tabelas a 10 pt com o cabeçalho a cinzento claro e linhas horizontais
   finas, tabela dos logótipos sem bordas e índice (campo TOC) atualizado ao abrir.
Uso: python scripts/gerar_docx.py 6   (gera relatorio/RELATORIO_v6.docx)
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[1]
REL = RAIZ / "relatorio"
REFERENCIA = REL / "formato" / "referencia.docx"
FONTE = "Times New Roman"
CINZENTO = "E0E0E0"  # cabeçalho das tabelas (cinzento claro, como no PDF)
AZUL_LINKS = RGBColor(0, 0, 255)  # cor dos links no PDF (urlcolor "blue" do pandoc)
# Leitura do markdown: listas logo a seguir a um parágrafo, sem linha em branco (como no PDF).
ENTRADA = "markdown+lists_without_preceding_blankline"


def pandoc() -> str:
    """Caminho do pandoc (no PATH ou na instalação por utilizador do Windows)."""
    local = Path(os.environ.get("LOCALAPPDATA", "")) / "Pandoc" / "pandoc.exe"
    return shutil.which("pandoc") or str(local)


def _fonte(estilo, tamanho: float | None = None, negrito: bool | None = None, italico: bool | None = None) -> None:
    """Fonte Times New Roman (também para texto asiático e complexo), tamanho, negrito e itálico, a preto."""
    f = estilo.font
    f.name = FONTE
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(att), FONTE)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(att), None)
    if tamanho:
        f.size = Pt(tamanho)
    if negrito is not None:
        f.bold = negrito
    if italico is not None:
        f.italic = italico
    f.color.rgb = RGBColor(0, 0, 0)


def _paragrafo(estilo, alinhamento=None, antes: float = 0, depois: float = 6, linhas: float = 1.15,
               esquerda: float | None = None, suspensa: float | None = None, com_seguinte: bool = False) -> None:
    pf = estilo.paragraph_format
    if alinhamento is not None:
        pf.alignment = alinhamento
    pf.space_before, pf.space_after, pf.line_spacing = Pt(antes), Pt(depois), linhas
    if esquerda is not None:
        pf.left_indent = Cm(esquerda)
    if suspensa is not None:
        pf.first_line_indent = Cm(-suspensa)
    pf.keep_with_next = com_seguinte


def _estilo(doc, nome: str):
    """Estilo de parágrafo com este nome (cria-o, baseado no Normal, se não existir)."""
    try:
        return doc.styles[nome]
    except KeyError:
        st = doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        return st


def _campo_pagina(paragrafo) -> None:
    """Acrescenta o campo PAGE (número de página) a um parágrafo."""
    for tipo, texto in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = OxmlElement("w:r")
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            r.append(fc)
        else:
            el = OxmlElement("w:instrText" if texto.strip() == "PAGE" else "w:t")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
            r.append(el)
        paragrafo._p.append(r)


def criar_referencia(destino: Path = REFERENCIA) -> Path:
    """Cria o modelo do Word (--reference-doc) a partir do modelo por omissão do pandoc."""
    destino.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run([pandoc(), "-o", str(destino), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(str(destino))
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
        s.footer.is_linked_to_previous = False
        p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _campo_pagina(p)
    J, C, E = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    _fonte(doc.styles["Normal"], 11)
    for nome in ("Body Text", "First Paragraph"):
        _fonte(_estilo(doc, nome), 11)
        _paragrafo(_estilo(doc, nome), J)
    _paragrafo(_estilo(doc, "Compact"), E, depois=2)
    for nome, tam, antes in (("Heading 1", 14, 18), ("Heading 2", 12, 12), ("Heading 3", 11, 10)):
        _fonte(doc.styles[nome], tam, negrito=True, italico=False)
        _paragrafo(doc.styles[nome], E, antes=antes, depois=6, linhas=1.0, com_seguinte=True)
    for nome in ("TOC Heading", "Title", "Subtitle", "Author", "Date", "Abstract", "Block Text", "Footnote Text",
                 "Table Caption", "Image Caption", "Hyperlink"):
        try:
            _fonte(doc.styles[nome])
        except KeyError:
            pass
    doc.styles["Hyperlink"].font.color.rgb = AZUL_LINKS  # links a azul, como no PDF
    for nivel in (1, 2):
        st = _estilo(doc, f"toc {nivel}")
        _fonte(st, 11, negrito=nivel == 1)
        _paragrafo(st, E, depois=2, linhas=1.0, esquerda=0.6 * (nivel - 1))
    proprios = {  # estilo: (tamanho, negrito, itálico, alinhamento, antes, depois)
        "Centrado": (11, False, True, C, 0, 0),
        "Titulo do relatorio": (14, True, False, E, 24, 12),
        "Autor": (11, False, False, E, 0, 2),
        "Afiliacao": (10, False, False, E, 0, 18),
        "Titulo sem indice": (14, True, False, E, 12, 6),
        "Resumo": (10, False, False, J, 0, 6),
        "Legenda": (11, False, False, J, 10, 4),
        "Referencia": (11, False, False, E, 0, 6),
    }
    for nome, (tam, neg, ita, al, antes, depois) in proprios.items():
        st = _estilo(doc, nome)
        _fonte(st, tam, negrito=neg, italico=ita)
        _paragrafo(st, al, antes=antes, depois=depois, com_seguinte=nome in ("Legenda", "Titulo sem indice"))
    _paragrafo(_estilo(doc, "Referencia"), E, depois=6, esquerda=1, suspensa=1)
    doc.save(str(destino))
    return destino


def _bordas(tabela, tipo: str | None) -> None:
    """Bordas da tabela: só linhas horizontais finas (tipo="single") ou nenhuma (tipo=None)."""
    tblpr = tabela._tbl.tblPr
    antigo = tblpr.find(qn("w:tblBorders"))
    if antigo is not None:
        tblpr.remove(antigo)
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        horizontal = lado in ("top", "bottom", "insideH")
        el.set(qn("w:val"), "single" if tipo and horizontal else "nil")
        if tipo and horizontal:
            el.set(qn("w:sz"), "6" if lado in ("top", "bottom") else "2")
            el.set(qn("w:color"), "000000")
        bordas.append(el)
    tblpr.append(bordas)


def _sombra(celula, cor: str) -> None:
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def ajustar_tabelas(doc) -> None:
    """Primeira tabela (logótipos): sem bordas. Restantes: 10 pt, cabeçalho cinzento, linhas horizontais finas."""
    for i, t in enumerate(doc.tables):
        if i == 0:
            _bordas(t, None)
            continue
        _bordas(t, "single")
        for j, linha in enumerate(t.rows):
            for c in linha.cells:
                if j == 0:
                    _sombra(c, CINZENTO)
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(10)


def atualizar_campos_ao_abrir(doc) -> None:
    """Pede ao Word para atualizar os campos (o índice) quando abre o documento."""
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def gerar(versao: str) -> Path:
    criar_referencia()
    destino = REL / f"RELATORIO_v{versao}.docx"
    subprocess.run([pandoc(), str(REL / "RELATORIO_word.md"), "-f", ENTRADA, "-o", str(destino), f"--reference-doc={REFERENCIA}",
                    f"--resource-path={REL}"], check=True, cwd=REL)
    doc = Document(str(destino))
    ajustar_tabelas(doc)
    atualizar_campos_ao_abrir(doc)
    doc.save(str(destino))
    return destino


if __name__ == "__main__":
    print(gerar(sys.argv[1] if len(sys.argv) > 1 else "0"))
