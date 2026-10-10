"""Testes de scripts/gerar_docx.py (o script carrega-se pelo caminho)."""
import importlib.util
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

CAMINHO = Path(__file__).resolve().parents[1] / "scripts" / "gerar_docx.py"
spec = importlib.util.spec_from_file_location("gerar_docx", CAMINHO)
gd = importlib.util.module_from_spec(spec)
spec.loader.exec_module(gd)


def _doc_com_tabelas():
    doc = Document()
    for _ in range(2):
        t = doc.add_table(rows=2, cols=2)
        for c in t._cells:
            c.paragraphs[0].add_run("x")
    return doc


def test_ajustar_tabelas_logotipos_sem_bordas_e_restantes_com_cabecalho_cinzento():
    doc = _doc_com_tabelas()
    gd.ajustar_tabelas(doc)
    logos, dados = doc.tables
    bordas = logos._tbl.tblPr.find(qn("w:tblBorders"))
    assert all(b.get(qn("w:val")) == "nil" for b in bordas)
    cab = dados.rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
    assert cab.get(qn("w:fill")) == gd.CINZENTO
    assert dados.rows[1].cells[0]._tc.tcPr is None or dados.rows[1].cells[0]._tc.tcPr.find(qn("w:shd")) is None
    assert dados.rows[1].cells[0].paragraphs[0].runs[0].font.size.pt == 10
    lados = {b.tag.split("}")[1]: b.get(qn("w:val")) for b in dados._tbl.tblPr.find(qn("w:tblBorders"))}
    assert lados["top"] == lados["insideH"] == "single" and lados["insideV"] == "nil"


def test_atualizar_campos_ao_abrir():
    doc = Document()
    gd.atualizar_campos_ao_abrir(doc)
    assert doc.settings.element.find(qn("w:updateFields")).get(qn("w:val")) == "true"


@pytest.mark.skipif(not Path(gd.pandoc()).exists(), reason="pandoc não instalado")
def test_criar_referencia(tmp_path):
    destino = gd.criar_referencia(tmp_path / "ref.docx")
    doc = Document(str(destino))
    assert doc.styles["Normal"].font.name == "Times New Roman" and doc.styles["Normal"].font.size.pt == 11
    assert doc.styles["Heading 1"].font.size.pt == 14 and doc.styles["Heading 1"].font.bold
    assert doc.styles["Heading 2"].font.size.pt == 12
    assert round(doc.sections[0].left_margin.cm, 1) == 2.0
    for nome in ("Centrado", "Titulo do relatorio", "Legenda", "Referencia", "Resumo"):
        assert doc.styles[nome] is not None
    assert "PAGE" in doc.sections[0].footer._element.xml
    assert doc.styles["Hyperlink"].font.color.rgb == gd.AZUL_LINKS
