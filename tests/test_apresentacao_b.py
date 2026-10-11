"""Testes de scripts/apresentacao_b.py (funções de dados e de imagem; sem PowerPoint)."""
import importlib.util
import sys
from dataclasses import replace
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from PIL import Image

RAIZ = Path(__file__).resolve().parents[1]


def _carregar(nome):
    spec = importlib.util.spec_from_file_location(nome, RAIZ / "scripts" / f"{nome}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[nome] = mod
    spec.loader.exec_module(mod)
    return mod


ab = _carregar("apresentacao_b")


@pytest.fixture(scope="module")
def numeros():
    return ab.ap.ler_numeros()


def test_confianca_classifica_os_cinco_dispositivos(numeros):
    c = {d: nivel for d, _, nivel, _ in ab.confianca(numeros, True, simples=True)}
    assert c == {"Haste femoral": "Sólida", "Stent coronário": "Muito próximos", "Scaffold ósseo": "Depende da carga",
                 "Par articular": "Indicativo", "Implante dentário": "Indicativo"}


def test_confianca_muda_com_os_numeros(numeros):
    fraca = replace(numeros, mc_haste={"TOPSIS": 9000, "WASPAS": 10000, "VIKOR": 10000},
                    dc_stent={"A": 0.05, "B": 0.03})
    c = {d: nivel for d, _, nivel, _ in ab.confianca(fraca, False)}
    assert c["Haste femoral"] == "Provável" and c["Stent coronário"] == "Provável"
    assert c["Scaffold ósseo"] == "Provável"


def test_funil_le_os_numeros_do_results():
    f = ab.funil()
    assert f["mp"] > f["nivel2"] > f["base"] > f["final"] == 5
    assert f["elementos"] == 118


def test_perfis_exemplo_tem_o_defeito_sem_carga_e_a_alergia_ao_titanio():
    ex = {d: (antes, depois) for d, _, antes, depois in ab.perfis_exemplo()}
    assert ex["Defeito sem carga"] == ("Compósito PCL/β-TCP (impressão 3D)", "β-TCP poroso")
    assert ex["Alergia ao titânio confirmada por teste"] == ("Ti cp grau 4", "Zircónia Y-TZP")


def test_intervalo_e_n_regras():
    lo, hi, u = ab.intervalo("Aço inox 316L", "Módulo de Young")
    assert (lo, hi, u) == (205.0, 210.0, "GPa")
    assert ab.n_regras() == len(pd.read_csv(RAIZ / "data" / "perfis_doente.csv"))


def test_nb_e_mil_nao_partem_os_numeros():
    assert ab.nb("10 000 a 9 999") == "10 000 a 9 999"
    assert " " not in ab.mil(10000)


def test_pasta_renders_escolhe_a_primeira_completa(tmp_path):
    a, b = tmp_path / "a", tmp_path / "b"
    a.mkdir(), b.mkdir()
    for c, _, _ in ab.DISPOSITIVOS:
        (b / f"{c}_transparente.png").write_bytes(b"")
    (a / "anca_transparente.png").write_bytes(b"")
    assert ab.pasta_renders([a, b]) == b
    with pytest.raises(FileNotFoundError):
        ab.pasta_renders([a])


def _render_teste(caminho):
    a = np.zeros((200, 300, 4), np.uint8)
    a[80:120, 120:180] = (200, 50, 50, 255)    # objeto opaco
    a[60:160, 20:290, 3] = np.maximum(a[60:160, 20:290, 3], 90)  # sombra larga e meio transparente
    Image.fromarray(a, "RGBA").save(caminho)


def test_recortar_transparente_segue_o_objeto_e_esbate_a_sombra(tmp_path):
    o = tmp_path / "x_transparente.png"
    _render_teste(o)
    d = ab.recortar_transparente(o, tmp_path / "cache")
    im = np.asarray(Image.open(d))
    assert im.shape[1] < 300 and im.shape[0] < 200      # recortado à volta do objeto
    assert im[..., 3].max() == 255                       # o objeto fica opaco
    assert im[0, :, 3].max() == 0 and im[:, 0, 3].max() == 0  # a sombra chega a zero nas margens


def test_logotipo_branco_mantem_a_transparencia(tmp_path):
    o = tmp_path / "logo.png"
    a = np.zeros((10, 10, 4), np.uint8)
    a[2:5, 2:5] = (10, 20, 30, 255)
    Image.fromarray(a, "RGBA").save(o)
    b = np.asarray(Image.open(ab.logotipo_branco(o, tmp_path)))
    assert (b[3, 3] == (255, 255, 255, 255)).all() and b[0, 0, 3] == 0


def test_imagem_reduzida(tmp_path):
    o = tmp_path / "grande.png"
    Image.new("RGB", (4000, 2000)).save(o)
    d = ab.imagem_reduzida(o, 1000, tmp_path / "c")
    assert Image.open(d).size == (1000, 500)


def test_confianca_sem_nomes_simples_mantem_a_versao_1(numeros):
    c = ab.confianca(numeros, True)
    assert c[0][1] == "Ti-6Al-4V ELI" and c[3][2] == "Indicação"


def test_razoes_fea_e_grafico(tmp_path):
    csv = tmp_path / "fea.csv"
    pd.DataFrame({"malha_mm": [1.0] * 4 + [2.0], "material": ["316L", "316L", "Ti-6Al-4V", "Ti-6Al-4V", "316L"],
                  "zona": [1, 7, 1, 7, 1], "razao": [0.6, 0.1, 0.7, 0.2, 0.9]}).to_csv(csv, index=False)
    r = ab.razoes_fea(csv)
    assert r == {"Ti-6Al-4V": (0.7, 0.2), "316L": (0.6, 0.1)}
    p = ab.grafico_fea(r, tmp_path / "g.png")
    assert Image.open(p).size[0] > 1000


def test_n_verificados_conta_os_estados_verificado(tmp_path):
    pd.DataFrame({"estado": ["Verificado", "Verificado (fornecedor)", "A verificar", "Verificado (derivado)"]}) \
        .to_csv(tmp_path / "materiais.csv", index=False)
    assert ab.n_verificados(tmp_path) == (3, 4)
    v, t = ab.n_verificados()
    assert 0 < v < t


def test_mc_primeiro_le_o_monte_carlo_das_propriedades(tmp_path):
    pd.DataFrame({"caso": ["Stent"] * 3 + ["Haste"], "analise": ["MC propriedades"] * 2 + ["W pesos ±20 %"] * 2,
                  "metodo": ["TOPSIS", "WASPAS", "TOPSIS", "TOPSIS"], "material": ["Pt-Cr"] * 4,
                  "pct_primeiro": [70.0, 56.0, 83.0, 1.0]}).to_csv(tmp_path / "monte_carlo.csv", index=False)
    assert ab.mc_primeiro("Stent", tmp_path) == {"Pt-Cr": {"TOPSIS": 70.0, "WASPAS": 56.0}}
    st = ab.mc_primeiro("Stent")
    assert set(st["Pt-Cr"]) == {"TOPSIS", "WASPAS", "VIKOR"}


def test_graficos_ficam_com_o_tamanho_da_caixa(tmp_path):
    """Os gráficos dos diapositivos têm exatamente as polegadas da caixa (300 ppp): a letra não encolhe."""
    p = ab.grafico_mc_stent({"Pt-Cr": {"TOPSIS": 70, "WASPAS": 56, "VIKOR": 58},
                             "Co-Cr L605": {"TOPSIS": 27, "WASPAS": 41, "VIKOR": 32}}, tmp_path / "m.png")
    assert Image.open(p).size == (round(5.55 * 300), round(4.7 * 300))
    g = ab.grafico_gibson(tmp_path / "g.png")
    assert Image.open(g).size == (round(5.55 * 300), round(4.7 * 300))
    f = ab.grafico_funil(pd.read_csv(RAIZ / "results" / "funil.csv"), tmp_path / "f.png")
    assert Image.open(f).size == (round(7.9 * 300), round(4.9 * 300))
    r = ab.grafico_fea({"316L": (0.5, 0.1)}, tmp_path / "r.png")
    assert Image.open(r).size == (round(5.25 * 300), round(4.3 * 300))


def test_dados_e_grafico_da_cfd(tmp_path):
    import json
    js = tmp_path / "cfd.json"
    js.write_text(json.dumps({"Pt-Cr": {"espessura_um": 79.0, "fracao_baixa": 0.46},
                              "316L": {"espessura_um": 135.0, "fracao_baixa": 0.81}}), encoding="utf-8")
    d = ab.dados_cfd(js)
    assert d["316L"]["fracao_baixa"] == 0.81
    p = ab.grafico_cfd(d, tmp_path / "c.png")
    assert Image.open(p).size[0] > 1000


def test_margem_leitura_literal_e_a_do_grafico_do_scaffold():
    """Os cartões do diapositivo do scaffold dizem a margem da mesma leitura que o gráfico (a literal)."""
    primeiro, segundo, dc = ab.margem_leitura_literal()
    assert primeiro == "β-TCP poroso" and segundo == "Vidro bioativo 45S5 poroso"
    assert 0.04 < dc < 0.05


def test_desafios_e_resposta_usam_o_problema_indicado_pelo_docente():
    desafios = {nome: d for _, nome, d in ab.DISPOSITIVOS}
    assert desafios["Par articular"] == "desgaste e partículas"
    assert desafios["Stent coronário"] == "resistência mecânica e biocompatibilidade"
    assert desafios["Scaffold ósseo"] == "suporte celular e degradação controlada"
    assert desafios["Implante dentário"] == "integração com o osso e corrosão"
    assert [r[0] for r in ab.RESPOSTA] == [nome for _, nome, _ in ab.DISPOSITIVOS]
    for linha in ab.RESPOSTA + ab.FABRICO:
        assert len(linha) == 4 and all(linha)
        assert not any(c in t for t in linha for c in (chr(0x2014), chr(0x2013)))  # sem travessões
    # recomendação do stent (decisão do André de 11 out): Pt-Cr e L605 praticamente equivalentes
    proposta = dict((r[0], r[3]) for r in ab.RESPOSTA)
    assert proposta["Stent coronário"] == (
        "Pt-Cr e L605 praticamente equivalentes: o Pt-Cr é favorecido pelo TOPSIS, pelo VIKOR e pela radiopacidade; o "
        "L605 pelo WASPAS e por ter todas as propriedades em fichas de fornecedor. Sem vencedor universal.")
    assert proposta["Scaffold ósseo"] == ("Depende da carga: compósito PCL/β-TCP impresso quando o defeito suporta "
                                          "carga (1.º no ranking direto em 5 das 6 combinações); β-TCP poroso quando "
                                          "não suporta (mais vezes 1.º na incerteza dos dados).")


def test_quadro_cabe_no_diapositivo_com_letra_minima():
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = ab.Inches(13.333), ab.Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    assert abs(sum(ab.LARG_RESPOSTA) - 12.13) < 1e-9
    cab = ab.quadro(s, ("Caso", "Problema", "Biomateriais", "Proposta"), ab.RESPOSTA, ab.LARG_RESPOSTA,
                    h_cab=0.4, **ab.ALT_RESPOSTA)
    assert len(cab) == 4
    fundo = max(Emu(f.top + f.height).inches for f in s.shapes)
    assert fundo < 6.95  # acima do rodapé
    tamanhos = [r.font.size.pt for f in s.shapes if f.has_text_frame for p in f.text_frame.paragraphs
                for r in p.runs]
    assert min(tamanhos) >= 14


def test_espectro_usa_as_versoes_curtas(numeros):
    """No espectro (diapositivo 9) o nome do caso está ao lado: a frase curta sem o prefixo, o resto igual."""
    st = {d: frase for d, _, _, frase in ab.confianca(numeros, True, simples=True)}
    assert ab.FRASE_STENT_CURTA == "Stent: Pt-Cr e L605 praticamente equivalentes"
    assert ab.FRASE_SCAFFOLD_CURTA == "Scaffold: depende da carga (compósito com carga; β-TCP sem carga)"
    assert st["Stent coronário"] == "Pt-Cr e L605 praticamente equivalentes"
    assert st["Scaffold ósseo"] == "depende da carga (compósito com carga; β-TCP sem carga)"


def test_sem_prefixo():
    assert ab.sem_prefixo("Stent: Pt-Cr e L605 praticamente equivalentes") == "Pt-Cr e L605 praticamente equivalentes"
    assert ab.sem_prefixo("sem prefixo") == "sem prefixo"


def test_sensibilidade_do_jovem_ativo_le_os_perfis_de_referencia():
    """H1: o TOPSIS passa ao Co-Cr-Mo (results/perfis_referencia.csv, cenário A); o Borda mantém o Ti-6Al-4V ELI."""
    desc, antes, depois, borda = ab.sensibilidade_perfil("H1")
    assert desc == "Jovem ativo"
    assert antes == "Ti-6Al-4V ELI" and depois == "Co-Cr-Mo forjado" and borda == "Ti-6Al-4V ELI"


def test_aviso_diz_que_nao_e_ferramenta_clinica():
    assert ab.AVISO == "Não é ferramenta clínica nem recomendação terapêutica."


def _fea_json(caminho, energias, com_malha):
    import json
    casos = {m: {"zonas": {"1": {"energia_J": e1}, "7": {"energia_J": e7}}} for m, (e1, e7) in energias.items()}
    if com_malha:
        casos = {"casos": {m: {"3.5": c} for m, c in casos.items()}}
    caminho.write_text(json.dumps(casos), encoding="utf-8")
    return caminho


def test_razoes_duas_ferramentas_e_diferenca_maxima(tmp_path):
    sw = _fea_json(tmp_path / "sw.json", {"Intacto": (2.0, 4.0), "Aço inox 316L": (1.0, 1.0)}, True)
    md = _fea_json(tmp_path / "md.json", {"Intacto": (2.0, 4.0), "Aço inox 316L": (1.02, 1.0)}, False)
    r = ab.razoes_duas_ferramentas(sw, md)
    assert r == {"Aço inox 316L": {"SOLIDWORKS": (0.5, 0.25), "Ansys": (0.51, 0.25)}}
    assert abs(ab.dif_max_razoes(r) - 2.0) < 1e-9


def test_comparacao_real_das_ferramentas_bate_com_o_documento():
    """Números de docs/ANSYS_COMPARACAO.md (ramo extensoes): diferenças de 2,3 % ou menos; 46,2 contra 45,8 %."""
    if not (ab.SW_JSON.exists() and ab.MAPDL_JSON.exists() and ab.FLUENT_JSON.exists()):
        pytest.skip("resultados das duas ferramentas ausentes")
    r = ab.razoes_duas_ferramentas()
    assert round(ab.dif_max_razoes(r), 1) <= 2.3
    ordem = sorted(r, key=lambda m: -r[m]["Ansys"][1])
    assert ordem == sorted(r, key=lambda m: -r[m]["SOLIDWORKS"][1]) == list(ab.NOMES_FEA_DUAS)
    a = ab.areas_duas_ferramentas()
    assert [round(a["Pt-Cr"][k], 1) for k in ("OpenFOAM", "Ansys Fluent")] == [46.2, 45.8]
    assert [round(a["316L"][k], 1) for k in ("OpenFOAM", "Ansys Fluent")] == [80.9, 79.8]


def test_grafico_duas_ferramentas_tem_o_tamanho_da_caixa(tmp_path):
    h = {"Aço inox 316L": {"SOLIDWORKS": (0.53, 0.27), "Ansys": (0.52, 0.27)}}
    s = {"Pt-Cr": {"OpenFOAM": 46.2, "Ansys Fluent": 45.8}}
    p = ab.grafico_duas_ferramentas(h, s, tmp_path / "d.png")
    assert Image.open(p).size == (round(12.13 * 300), round(3.6 * 300))


def test_primeiro_existente(tmp_path):
    a, b = tmp_path / "a.png", tmp_path / "b.png"
    assert ab.primeiro_existente([a, b]) is None
    b.write_bytes(b"")
    assert ab.primeiro_existente([a, b]) == b


def test_capa_video_tira_a_ultima_imagem(tmp_path):
    import imageio_ffmpeg
    mp4 = tmp_path / "v.mp4"
    w = imageio_ffmpeg.write_frames(str(mp4), (64, 48), fps=10, macro_block_size=16)
    w.send(None)
    for k in range(10):
        w.send(np.full((48, 64, 3), 25 * k, np.uint8).tobytes())
    w.close()
    im = np.asarray(Image.open(ab.capa_video(mp4, tmp_path / "capa.png")).convert("L"))
    assert im.shape == (48, 64) and im.mean() > 180  # a última imagem é a mais clara


def test_palavras_notas_e_minutos(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    for texto in ("um dois três", "quatro cinco", "reserva não conta"):
        prs.slides.add_slide(prs.slide_layouts[6]).notes_slide.notes_text_frame.text = texto
    assert ab.palavras_notas(prs, 2) == 5
    assert ab.minutos_fala(1300, 15) == pytest.approx(10.25)


def test_vencedores_ranking_do_scaffold_e_a_contagem_da_frase():
    """A contagem "5 das 6 combinações" da frase do scaffold vem de results/pontuacoes.csv."""
    v = ab.vencedores_ranking("Scaffold")
    assert len(v) == 6 and ab.n_primeiro(v, ab.COMPOSITO) == 5
    assert v[("A", "TOPSIS")] == ab.BTCP
    assert "5 das 6 combinações" in ab.FRASE_SCAFFOLD


def test_espectro_do_scaffold_diz_a_carga(numeros):
    sc = [linha for linha in ab.confianca(numeros, True, simples=True) if linha[0] == "Scaffold ósseo"][0]
    assert sc[2] == "Depende da carga" and "carga" in sc[3]


def test_grafico_scaffold_tem_o_tamanho_da_caixa(tmp_path):
    v = {(c, m): ab.COMPOSITO for c in "AB" for m in ab.METODOS}
    pct = {ab.BTCP: {m: 60.0 for m in ab.METODOS}, ab.COMPOSITO: {m: 20.0 for m in ab.METODOS}}
    p = ab.grafico_scaffold(v, pct, tmp_path / "s.png")
    assert Image.open(p).size == (round(8.6 * 300), round(4.84 * 300))


def test_quadro_aceita_alturas_por_linha():
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ab.quadro(s, ("a", "b"), [("x", "y"), ("z", "w")], (1.0, 1.0), y0=1.0, h_cab=0.4, h_linha=0.5, alturas={1: 1.0})
    cartoes = sorted((f for f in s.shapes if not f.has_text_frame or not f.text_frame.text), key=lambda f: f.top)
    assert [round(Emu(c.height).inches, 2) for c in cartoes][-2:] == [0.5, 1.0]
